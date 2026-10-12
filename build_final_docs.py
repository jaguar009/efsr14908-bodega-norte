"""Build the final course documents from verified artifacts, without private credentials."""
from pathlib import Path
import json,re,shutil
from docx import Document
from docx.shared import Cm,Pt,RGBColor
import build_word_docs as b

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'docs/word';OUT.mkdir(parents=True,exist_ok=True)
EVID=ROOT/'docs/evidencias/final_2026-10-11';EVID.mkdir(parents=True,exist_ok=True)
STATUS='Validación final al 11 de octubre de 2026: 25 pruebas locales, 17 casos de backend HTTP con PostgreSQL aislado, 11 recorridos de interfaz y 8 consultas contra IBM real aprobados. El programa Java se ejecutó contra el backend de prueba. Las respuestas de Supabase Auth e IBM de los 17 casos son fixtures locales; las 8 consultas IBM tienen registro independiente de conexión real. La base principal no cambió durante las pruebas. La aceptación por personas del Grupo 10 o una bodega real y las horas de asistencia no se certifican con estas pruebas.'

def new(name,subtitle=None):
    d=Document();b.cover(d,name,subtitle);return d
def write(d,name):
    for p in d.paragraphs:
        if p.text.startswith('Fecha de elaboración:'):p.text='Fecha de elaboración: 12 de octubre de 2026'
        if 'Semestre no consignado en los documentos recibidos' in p.text:p.text=p.text.replace('Semestre no consignado en los documentos recibidos','Semestre por confirmar')
        if p.style.name.startswith('Heading'):
            p.text=re.sub(r'^\d+(?:\.\d+)*\.?'+' '+r'', '', p.text)
    for p in all_paragraphs(d):
        if 'Gonzalo Yoel Choque Becerra' in p.text:p.text=p.text.replace('Gonzalo Yoel Choque Becerra','Gonsalo Yoel Choque Becerra')
    d.core_properties.author='Grupo 10';d.core_properties.title=name.replace('_',' ').replace('.docx','')
    b.save(d,name)
def text(d,s):b.para(d,s)
def section(d,s):b.h(d,s)
def steps(d,items):b.numbered(d,items)
def grid(d,headers,rows,widths=None):b.table(d,headers,rows,widths)

# Curated real captures from fresh isolated browser contexts.
ui=json.loads((ROOT/'output/playwright-final/validacion_interfaz.json').read_text(encoding='utf-8'))
assert ui['allPassed'] and len(ui['checks'])==11 and not ui['errors'] and not ui['warnings']
for shot in ui['screenshots']:shutil.copy2(ROOT/'output/playwright-final'/shot['file'],EVID/shot['file'])
shutil.copy2(ROOT/'output/playwright-final/validacion_interfaz.json',EVID/'validacion_interfaz.json')
backend=json.loads((ROOT/'docs/evidencias/validacion_final_2026-10-11.json').read_text(encoding='utf-8'))
ibm=json.loads((ROOT/'docs/evidencias/validacion_watson_2026-10-11.json').read_text(encoding='utf-8'))
assert backend['allPassed'] and backend['productionUnchanged'] and len(backend['cases'])==17
assert ibm['allPassed'] and len(ibm['cases'])==8
faq=json.loads((ROOT/'integration/watson/preguntas_frecuentes.json').read_text(encoding='utf-8'))

# Preserve the course report's diagnostic, objectives, Canvas and plan.
captured={};real_save=b.save;b.save=lambda d,n:captured.update({n:d})
b.doc_03();b.doc_04();b.save=real_save
d=captured['03_Entrega_Final_Informe.docx']
replacements={
 'backend ASP.NET Core MVC preparado para Supabase PostgreSQL':'backend ASP.NET Core MVC publicado y conectado a Supabase PostgreSQL',
 'en un entorno demostrativo local.':'en un entorno académico, con demo local y publicación compartida.',
 'posible despliegue futuro sujeto a configuración y seguridad':'GitHub Pages conectado al backend MVC en Render y Supabase',
 'Uso guiado por módulos, manual y ayuda local de demostración.':'Uso guiado por módulos, manual, ayuda local y Watson Assistant en español.',
 'dentro de una transacción serializable.':'dentro de una transacción ReadCommitted con bloqueos de filas y control de reintentos.',
 'Copias automatizadas y despliegue operativo aún no verificado. Autenticación, roles y auditoría sí se implementaron.':'Copias programadas automáticas y operación real multiempresa. Se entregan respaldo manual y herramienta de recuperación en un esquema nuevo.',
 'Conexión productiva con IBM Watson Assistant. Se entrega la especificación; la interfaz actual mantiene respuestas locales de demostración.':'Uso permanente de IBM más allá del Trial gratuito de 30 días. La conexión real de ayuda está implementada; la ayuda local funciona de forma independiente.',
 'Integración Watson sin credenciales o servicio configurado':'Vencimiento del Trial de IBM o agotamiento de su cuota',
 'Mantener el contrato documentado y etiquetar la conexión como pendiente.':'Conservar la ayuda local y el contenido exportable; no contratar planes de pago sin una decisión del titular.',
 'las operaciones autenticadas requieren validación':'las operaciones de negocio fueron probadas en un esquema aislado con Auth simulado',
 'Demo local; Watson no está conectado.':'Ayuda local y ocho acciones Watson verificadas contra IBM real.',
 'Respuestas offline en interfaz. Watson Assistant se describe como integración pendiente.':'Ayuda local de stock y ventas; Watson Assistant explica ocho tareas operativas en español mediante el backend autenticado.',
 'Servicio conversacional de IBM considerado para una integración futura.':'Servicio conversacional de IBM integrado mediante Assistant builder de watsonx Orchestrate; Trial gratuito limitado a 30 días.',
 'Manual de usuario con los pasos para venta, inventario, productos y reportes.':'Manual de usuario con venta, inventario, anulaciones, permisos, reportes, respaldo y asistentes.',
 'Manual de usuario", "Flujos de consulta, productos, venta, inventario, reportes y asistente local.':'Manual de usuario", "Flujos de consulta, productos, venta, inventario, reportes y asistentes.',
 'Demo local; Watson no está conectado.':'Ayuda local y conexión real Watson verificadas.',
 'HomeController, ProductsController, SalesController, HealthController, BodegaService y BodegaRepository.':'Controladores MVC, BodegaService, BodegaRepository, WatsonAssistantService y herramientas BodegaNorte.Tools.',
 'publicación de la demo estática.':'publicación del frontend conectado a la API de Render.',
 'construir la interfaz con npm.':'construir la interfaz con pnpm y el lockfile.',
 'Implementado en modo SQL; ver resultados técnicos en la matriz.':'Implementado y verificado mediante HTTP y PostgreSQL en esquema aislado.',
 'Implementado en la interfaz de demostración.':'Implementado; filtro y actualización de existencias verificados.',
 'Incluido en Reportes; pendiente de aceptación manual del equipo.':'Descarga CSV y contenido verificados en navegador.',
 'Funcionalidad incluida; la aceptación de interfaz queda pendiente de registro del equipo.':'Catálogo de 12 productos demo incluido; alta, edición y archivo verificados.',
 'Filtro y exportación incluidos; la validación visual del equipo queda pendiente.':'Filtro, CSV y vista móvil verificados técnicamente.',
 'Se adjuntan fuera de este informe para conservar el documento en blanco y negro.':'Las capturas a color se incorporan en Anexo_Validacion_Final_Grupo10.docx.',
 'Antes de operar, validar autenticación, roles, auditoría y recuperación del respaldo en el entorno publicado; proteger credenciales y gestionar los datos personales.':'Mantener copias privadas, ensayar la recuperación en un destino nuevo y revisar la vigencia del Trial de IBM; conservar la ayuda local y proteger las credenciales.',
 '[[FN1]]':'(PRODUCE, Las Mipyme en cifras 2024).','[[FN2]]':'(ANPD, Nuevo Reglamento de Protección de Datos Personales).',
 '[[FN3]]':'(INEI, TIC en las empresas 2023).','[[FN4]]':'(MINAM, Decreto Supremo 009-2019-MINAM).',
 'database/README.md':'06_Guia_Instalacion.docx',
 'Registro de npm run build.':'Registro de pnpm run build:github y build:visualstudio.',
 '/campanas/':'/campa%C3%B1as/',
 'docs/cronograma.md, docs/gestion/, docs/matriz_pruebas.md':'04_Cronograma.docx, 11_Backlog_Priorizado.docx, 12_Plan_Integracion_Continua.docx y 08_Matriz_Pruebas.docx',
}
def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:yield from cell.paragraphs
for p in all_paragraphs(d):
    s=p.text
    for a,z in replacements.items():s=s.replace(a,z)
    if s.startswith('Validación técnica: 18 pruebas'):s=STATUS+' Evidencia detallada en Anexo_Validacion_Final_Grupo10.docx.'
    elif s=='Pendiente para Supabase PostgreSQL':s='Aprobado en PostgreSQL aislado'
    elif s.startswith('La prueba aprobada del paquete anterior fue sobre SQL Server') or s.startswith('La prueba histórica SQL Server'):s='Registro actual de BodegaNorte.Tools y ejecución Java HTTP en PostgreSQL; ver Anexo_Validacion_Final_Grupo10.docx.'
    elif s=='Pendiente de aceptación manual del equipo':s='Validación técnica de navegador aprobada'
    elif s.startswith('Capturas en docs/evidencias/ como referencia'):s='Once recorridos de navegador, CSV descargado y revisión móvil; capturas a color en el anexo.'
    elif s.startswith('Capturas de referencia:'):s='Anexo_Validacion_Final_Grupo10.docx contiene las capturas actuales a color y los resultados de pruebas. Los archivos originales se adjuntan en Evidencias.'
    elif s.startswith('Se implementaron Supabase Auth y permisos'):s+=' Se añadió recuperación validada y conexión real Watson Assistant con límites de consulta.'
    elif s.startswith('El backend Render está publicado y conectado'):s='El frontend público está en https://jaguar009.github.io/efsr14908-bodega-norte/ y el backend MVC en https://bodega-norte-grupo10.onrender.com. Supabase y las credenciales IBM se configuran de forma privada. El paquete no contiene contraseñas, firmas ficticias ni registros de clientes reales.'
    elif s=='Actualización operativa al 04-10-2026':s='Validación final y publicación'
    if s!=p.text:p.text=s
for p in list(d.paragraphs):
    if p.text.startswith('3. Diagnóstico'):p.insert_paragraph_before('Capítulo I Diagnóstico del caso',style='Heading 1')
    elif p.text.startswith('4. Objetivos'):p.insert_paragraph_before('Capítulo II Propuesta y alcance',style='Heading 1')
    elif p.text.startswith('7. Productos'):p.insert_paragraph_before('Capítulo III Gestión producto y recursos',style='Heading 1')
section(d,'Uso de Watson Assistant')
text(d,'Se creó la instancia Bodega Norte Grupo 10 en IBM Cloud, región Sídney, plan Trial gratuito. Assistant builder conserva ocho acciones en español y 32 ejemplos. La versión 1 se publicó en el entorno Live. El servidor intercambia la clave privada por un token IAM temporal y llama a la API v2; no envía registros de la base ni ejecuta operaciones a partir de las respuestas.')
text(d,'En Asistente, el usuario elige ayuda local o IBM Watson. La ayuda local calcula stock bajo y ventas del día. IBM explica cómo utilizar los módulos. Al vencer el Trial o fallar IBM, las funciones de inventario y venta siguen disponibles y se conserva la ayuda local. La instancia gratuita no se presenta como un servicio permanente.')
grid(d,['Validación final','Resultado y alcance'],[
 ('25 pruebas locales','Reglas de dominio, reintentos, anulaciones, CSV y recuperación demo aprobadas.'),
 ('17 casos HTTP y PostgreSQL','Servidor real local, esquema desechable, transacciones reales; Auth e IBM simulados para aislamiento.'),
 ('11 recorridos de interfaz','Datos ficticios en contexto limpio, interacción real, descargas y vista móvil; cero errores y advertencias.'),
 ('8 consultas IBM reales','Las ocho respuestas publicadas coinciden con el contenido esperado.'),
 ('Publicación','GitHub Pages y Render actualizados; readiness 200, rutas privadas 401 sin sesión, CORS permitido.'),
 ('Asistencia y aceptación','Actividades propuestas; las fechas, horas, firmas y aceptación de negocio corresponden al equipo y a la docente.')],[4,12])
text(d,'Bibliografía técnica complementaria: IBM, Assistant v2 API, https://cloud.ibm.com/apidocs/assistant-v2; IBM, acceso al Trial de watsonx Orchestrate, https://www.ibm.com/docs/en/watsonx/watson-orchestrate/base?topic=orchestrate-accessing-watsonx-trial-version; Supabase, Auth getUser, https://supabase.com/docs/reference/javascript/auth-getuser. Consulta del 11 de octubre de 2026.')
# Place current product evidence in Chapter III and its sources in Bibliography.
body=d._element.body
conclusion=next(p._p for p in d.paragraphs if p.text=='11. Conclusiones')
final_start=next(p._p for p in d.paragraphs if p.text=='Validación final y publicación')
bibliography=d.paragraphs[-1]._p
tail=list(body)[list(body).index(final_start):]
for element in tail:
    if element is bibliography or element.tag.endswith('sectPr'):continue
    body.remove(element);conclusion.addprevious(element)
annex_heading=next(p._p for p in d.paragraphs if p.text=='15. Anexos')
body.remove(bibliography);annex_heading.addprevious(bibliography)
write(d,'03_Entrega_Final_Informe.docx')
write(captured['04_Cronograma.docx'],'04_Cronograma.docx')

d=new('Manual de usuario Bodega Norte','Operación del sistema y sus permisos')
text(d,'El sistema organiza inventario y ventas del caso académico. La web compartida requiere una cuenta confirmada y habilitada. La demo guarda datos ficticios en el navegador y simula roles. Ninguno de sus medios de pago realiza transferencias bancarias.')
sections=[('Ingreso y recuperación',['Abra la web publicada e inicie sesión con su cuenta confirmada.','Si olvidó la contraseña, use Olvidé mi contraseña, escriba el correo correcto y abra el enlace recibido; revise spam.','Un administrador debe habilitar las otras cuentas ya creadas y confirmadas.']),('Venta',['Abra Nueva venta, busque productos y agréguelos al carrito.','Ajuste cantidades y compruebe el total; no supere el stock.','Pulse Cobrar venta, elija el medio de pago y confirme Registrar venta.','Revise la constancia interna. Si la conexión falla, reintente el mismo cobro pendiente para evitar duplicados.']),('Productos y existencias',['El administrador crea, edita o archiva productos desde Productos o Inventario.','Use el filtro Bajo para existencias menores o iguales al mínimo.','Para cambiar stock, use Ajustar stock con cantidad positiva o negativa y motivo. Si aparece un conflicto de versión, actualice los datos.']),('Anulaciones y movimientos',['El administrador abre Historial de ventas y anula una venta con motivo.','La venta queda en el historial como anulada y devuelve existencias una vez.','Movimientos de stock permite revisar ingresos, salidas, motivos y responsables.']),('Proveedores y reportes',['El administrador registra contactos de proveedores. No elimine un proveedor con productos activos vinculados.','En Reportes seleccione fechas; se usan fechas de Lima y se excluyen anulaciones.','Exportar CSV descarga los datos mostrados; un periodo vacío muestra cero.']),('Configuración y permisos',['El administrador cambia datos de la bodega, alertas y permisos de cuentas confirmadas.','El cajero consulta y registra ventas; no gestiona productos, anulaciones ni respaldos.','La auditoría registra operaciones administrativas. El modo oscuro se guarda en el navegador.']),('Respaldo y recuperación',['En Configuración y Respaldo de datos, el administrador descarga una copia JSON privada.','La demo valida el archivo y exige escribir RESTAURAR antes de reemplazar los datos locales.','El respaldo PostgreSQL se recupera mediante BodegaNorte.Tools en un esquema nuevo; consulte la guía de instalación.']),('Asistentes',['Abra Asistente y elija Ayuda local para stock bajo o ventas de hoy.','IBM Watson responde preguntas de uso, por ejemplo Cómo registro una venta. Su consulta se envía a IBM; no incluya datos personales.','Las respuestas orientan y no ejecutan operaciones. La ayuda local sigue disponible si IBM falla o termina el Trial.'])]
for title,items in sections:section(d,title);steps(d,items)
write(d,'05_Manual_Usuario.docx')

d=new('Guía de instalación Bodega Norte','Visual Studio MVC Supabase y recuperación')
section(d,'Requisitos y estructura');text(d,'Instale Visual Studio con desarrollo ASP.NET y el SDK .NET 8, Node.js y pnpm. Java 17 o superior solo se requiere para la prueba Java. Extraiga todo el ZIP; la solución está en 06_Codigo/EFSR14908_Bodega_Norte/backend/BodegaNorte.sln.')
section(d,'Demostración local');steps(d,['Abra una terminal en la raíz del código y ejecute pnpm install --frozen-lockfile.','Ejecute pnpm run dev y abra la dirección que muestre Vite. La demo trabaja con datos ficticios locales.'])
section(d,'Visual Studio con base compartida');steps(d,['Ejecute pnpm run build:visualstudio en la raíz del código.','Abra backend/BodegaNorte.sln; el proyecto de inicio es BodegaNorte.Api.','En el proyecto API, abra Manage User Secrets y configure los valores privados de la tabla. En este equipo ya están guardados; el ZIP no los incluye.','En una base nueva aplique schema.sql, upgrade_operations.sql y seed.sql, en ese orden. No vuelva a sembrar una base operativa.','Ejecute con F5 y visite https://localhost:7248 o http://localhost:5248. /api/health/ready confirma conexión y tablas.'])
grid(d,['Configuración privada','Valor o procedencia'],[('ConnectionStrings:BodegaNorte','Cadena Npgsql del Session pooler Supabase con contraseña real y TLS VerifyFull.'),('Supabase:Url','URL del proyecto Supabase conectado.'),('Supabase:PublishableKey','Clave pública del proyecto; no es una service role.'),('BootstrapAdminEmail','Correo del administrador inicial confirmado. Las demás cuentas se habilitan desde Configuración.'),('Watson:Enabled','true después de configurar y verificar el asistente.'),('Watson:ServiceUrl','Service instance URL del asistente en Assistant settings y API details; puede diferir de la URL general de Orchestrate.'),('Watson:AssistantId y Watson:EnvironmentId','IDs del asistente y entorno Live.'),('Watson:ApiKey','Clave privada IBM; solo en User Secrets o entorno del backend.')],[6,10])
section(d,'Pruebas aisladas');text(d,'Desde la raíz del código: dotnet run --project backend/BodegaNorte.Tools -c Release -- verify --report output/validacion_final.json. El comando crea esquemas desechables y ejecuta la API local con Auth e IBM simulados; no inserta cuentas ni ventas en la base principal. Necesita los secretos privados y Java disponible en PATH o en la ruta documentada del equipo de validación.')
text(d,'Para verificar IBM real: dotnet run --project backend/BodegaNorte.Tools -c Release -- assistant --report output/validacion_ibm.json. Ejecuta las ocho preguntas del archivo preguntas_frecuentes.json sin enviar datos de la base.')
section(d,'Recuperar PostgreSQL');steps(d,['Guarde el respaldo JSON privado descargado por el administrador.','Ejecute dotnet run --project backend/BodegaNorte.Tools -c Release -- restore "ruta/al/respaldo.json" --schema bodega_recovery_revision.','El destino debe ser nuevo y empezar por bodega_recovery_. El comando rechaza bodega_norte y esquemas existentes.','Compruebe los resultados y conteos. La herramienta no recrea contraseñas, Auth ni membresías; habilite los accesos mediante el procedimiento normal.','Configure Database:Schema con el nuevo destino solo después de su revisión. Conserve el esquema original durante la recuperación.'])
section(d,'Publicación');text(d,'GitHub Pages aloja React; Render ejecuta ASP.NET Core MVC; Supabase conserva PostgreSQL y Auth. Los workflows compilan y publican desde master. En Render se usa __ para las claves anidadas, por ejemplo Watson__ApiKey. Las credenciales permanecen fuera del repositorio.')
write(d,'06_Guia_Instalacion.docx')

d=new('Manual técnico Bodega Norte','Arquitectura contratos y reglas de negocio')
section(d,'Arquitectura');text(d,'La vista React usa Supabase Auth para obtener una sesión. La API MVC valida el usuario en Supabase, exige correo confirmado y consulta members para obtener el rol. BodegaService aplica reglas y BodegaRepository usa Npgsql. El esquema privado bodega_norte contiene nueve tablas sin acceso directo para anon o authenticated.')
b.image(d,'docs/diagramas/arquitectura_bw.png','Arquitectura MVC y Supabase',6.1)
section(d,'Datos');grid(d,['Tabla','Responsabilidad'],[(x,y) for x,y in [('categories','Clasificación de productos.'),('suppliers','Contactos de proveedores.'),('products','Stock, mínimo, costo, precio y versión.'),('sales','Venta, estado, pago e identificador de reintento.'),('sale_items','Productos y precios de la venta.'),('members','Rol y acceso de identidades Supabase Auth.'),('settings','Datos de la bodega y alertas.'),('stock_movements','Cambios de stock con motivo y saldo.'),('audit_log','Trazabilidad de operaciones administrativas.')]], [4,12])
section(d,'Reglas y contratos');b.bullets(d,['La venta usa requestId y bloqueos para confirmar una sola vez el mismo cobro; reutilizar la clave con otro contenido se rechaza.','El servidor valida disponibilidad, cantidades y precio. Las operaciones se confirman o revierten en transacciones ReadCommitted con bloqueos de filas.','Las actualizaciones de existencias comprueban expectedVersion. Las anulaciones son idempotentes y mantienen el límite numérico del stock.','El rol proviene de members; user_metadata no habilita privilegios. Un cajero no accede a rutas administrativas.','El respaldo PostgreSQL es formato postgresql-v2; la demo usa demo-v2. No se importan formatos cruzados.','Los reportes usan America/Lima y excluyen anulaciones; no inventan costos de ventas históricas.','POST /api/assistant/message exige Staff, máximo 500 caracteres y 20 consultas por minuto por usuario. GET /api/assistant/status informa configuración sin exponer la clave.'])
section(d,'Recuperación y operación');text(d,'BodegaNorte.Tools recupera únicamente tablas autorizadas en un esquema nuevo, con validación y rollback completo ante errores. No restaura miembros ni usuarios Auth. La readiness comprueba las nueve tablas y las columnas de operación. Supabase exige validación TLS del servidor y certificado público incluido.')
text(d,STATUS);write(d,'07_Manual_Tecnico.docx')

d=new('Matriz de pruebas Bodega Norte','Resultados técnicos y alcance de la validación')
text(d,STATUS)
grid(d,['Suite','Casos','Entorno y resultado'],[('Dominio JavaScript','25','Aprobados con node --test; lógica local y demo.'),('HTTP y PostgreSQL','17','Aprobados con Kestrel y PostgreSQL aislado; Auth e IBM son fixtures.'),('Interfaz','11','Aprobados en Chrome limpio y datos demo; sin errores ni advertencias.'),('IBM real','8','Aprobados en Assistant builder Live del Trial.')],[4,2,10])
section(d,'Casos HTTP y PostgreSQL');grid(d,['ID','Caso','Resultado'],[(x['id'],x['description'],'Aprobado') for x in backend['cases']],[2.2,11,2.8])
section(d,'Recorridos de interfaz');grid(d,['ID','Caso','Resultado'],[(x['id'],x['description'],'Aprobado') for x in ui['checks']],[2.2,11,2.8])
section(d,'Consultas IBM reales');grid(d,['Pregunta','Resultado'],[(x['question'],'Respuesta esperada recibida') for x in ibm['cases']],[10,6])
section(d,'Registros entregados');text(d,'Los JSON de validación del backend, interfaz, IBM y despliegue se adjuntan en Evidencias. El anexo Word muestra resultados y capturas a color. Estas pruebas técnicas no sustituyen actas de aceptación, asistencia, horas ni una medición del impacto en una empresa real.')
write(d,'08_Matriz_Pruebas.docx')

d=new('Guion de sustentación Bodega Norte','Presentación del proyecto del Grupo 10')
section(d,'Distribución propuesta');grid(d,['Integrante','Contenido para explicar'],[(b.COORDINATOR,'Problema ficticio, objetivos, Canvas y plan.'),(b.TEAM[0],'Interfaz, venta, inventario, reportes y móvil.'),(b.TEAM[1],'Modelo PostgreSQL, permisos de esquema y recuperación.'),(b.TEAM[2],'Backend MVC, Auth, transacciones, pruebas e integración IBM.')],[8,8])
section(d,'Demostración');steps(d,['Explique que el caso usa datos académicos y que el cronograma y horas se ratifican con la docente.','Muestre una venta y el descuento de stock en un entorno de demostración.','Muestre rechazo por stock insuficiente, anulación y movimientos.','Abra reportes y exporte CSV; explique fecha de Lima y ventas anuladas.','Compare los permisos simulados de administrador y cajero; explique el control real del servidor.','Muestre recuperación demo y explique que PostgreSQL se restaura en un esquema nuevo.','Abra Watson y pregunte Cómo registro una venta; identifique el proveedor y la vigencia del Trial.','Presente la matriz y el anexo. Distinga las pruebas aisladas de la conexión real IBM y de la aceptación humana.'])
section(d,'Cierre');text(d,'El producto implementa el alcance académico y está publicado con sus manuales. El siguiente paso para uso comercial es validar el flujo con usuarios reales, gestionar datos personales y acordar mantenimiento y continuidad de servicios.')
write(d,'09_Guion_Sustentacion.docx')

d=new('Integración Watson Assistant Bodega Norte','Configuración funcionamiento y evidencias de IBM')
section(d,'Implementación');text(d,'Watson Assistant se utiliza mediante Assistant builder de watsonx Orchestrate. La instancia del Grupo 10 está en Sídney y usa Trial gratuito de 30 días. La versión 1 está en Live; las ocho respuestas se verificaron con la API v2 real. No se adquirió un plan de pago.')
grid(d,['Consulta','Ejemplos y propósito'],[(x['question'],x['examples'][0]) for x in faq],[8,8])
section(d,'Flujo y configuración');steps(d,['La interfaz ofrece IBM solo cuando GET /api/assistant/status indica configured true.','El usuario elige IBM y envía una consulta de hasta 500 caracteres.','La API valida la sesión Supabase y el rol Staff, limita la cuota y usa IAM para obtener un token IBM temporal.','El servidor llama al asistente y al entorno Live configurados; la clave nunca se envía al navegador.','La respuesta orienta sobre el uso del sistema; no ejecuta operaciones ni recibe el catálogo o las ventas.'])
text(d,'User Secrets: Watson:Enabled, Watson:ServiceUrl, Watson:AssistantId, Watson:EnvironmentId y Watson:ApiKey. En Render se usan los mismos nombres con __. Obtenga la URL específica en Assistant settings, Assistant IDs and API details; no la confunda con la URL general de Orchestrate.')
section(d,'Contenido y reproducción');text(d,'integration/watson/bodega_norte_acciones.json conserva el workspace de acciones sin credenciales. preguntas_frecuentes.json contiene ocho preguntas, respuestas y 32 ejemplos. El endpoint oficial Update skill permite importar el workspace a la habilidad de acciones; cree una versión y publíquela en Live después de entrenarla y comprobarla.')
section(d,'Validación y continuidad');text(d,'La herramienta assistant de BodegaNorte.Tools ejecutó las ocho consultas y recibió sus respuestas esperadas. Los casos IBM del verificador aislado comprueban el contrato, configuración y cuota con un fixture; no se presentan como conexión real. El Trial vence tras 30 días. IBM requiere contratar otro plan para continuar; el proyecto mantiene ayuda local y no cambia automáticamente el plan.')
text(d,'Fuente: IBM Assistant v2 API, https://cloud.ibm.com/apidocs/assistant-v2; Trial, https://www.ibm.com/docs/en/watsonx/watson-orchestrate/base?topic=orchestrate-accessing-watsonx-trial-version.')
write(d,'10_Integracion_Watson_Assistant.docx')

d=new('Backlog y criterios de aceptación Bodega Norte','Estado del alcance académico al cierre técnico')
rows=[('Catálogo','Alta, edición, archivo y búsqueda de productos.','DATA-01 y UI-03'),('Venta','Una venta por requestId y stock suficiente.','SALE-01, SALE-02 y SALE-03'),('Inventario','Ajuste con motivo y conflicto de versión.','STOCK-01 y UI-05'),('Permisos','Correo confirmado y rol desde members.','AUTH-01, AUTH-02 y UI-09'),('Reportes','Periodos de Lima, anulaciones excluidas y CSV.','Pruebas de dominio y UI-06'),('Respaldo','Formatos validados y recuperación en destino nuevo.','BACKUP-01, BACKUP-02 y UI-08'),('Watson','Ocho acciones y credenciales solo en servidor.','IBM-01 a IBM-03 y ocho consultas reales'),('Interfaz móvil','Menú usable y ausencia de desbordamiento.','UI-10')]
grid(d,['Funcionalidad','Criterio','Verificación'],rows,[3,8,5]);section(d,'Control del alcance');text(d,'Los criterios técnicos están aprobados. La asignación de roles del equipo procede del Entregable 1; no acredita contribuciones individuales ni reuniones Scrum. Las actas de aceptación y horas reales se completan únicamente con información del equipo. Compras, SUNAT, contabilidad y operación multiempresa no forman parte del alcance.')
write(d,'11_Backlog_Priorizado.docx')

d=new('Plan de integración continua Bodega Norte','Compilación publicación y recuperación de versiones')
steps(d,['Mantenga cambios en Git y revise el diff antes de publicar. Las credenciales quedan en User Secrets o variables privadas.','El workflow ci.yml ejecuta las pruebas de dominio, construye GitHub Pages y Visual Studio, compila .NET y la prueba Java.','deploy-pages.yml publica React desde master en GitHub Pages. Render despliega automáticamente el backend desde el mismo repositorio.','La validación con PostgreSQL e IBM real se ejecuta aparte con secretos privados; no se introducen credenciales de producción en las pruebas públicas de CI.','Después del despliegue compruebe readiness, rechazo anónimo, CORS y los flujos autorizados que correspondan.','Para recuperar una versión use un commit revisado y despliegue la versión elegida; conserve copias de datos privadas y recupere en un esquema nuevo.'])
text(d,'La versión b31a431 se publicó con ambos workflows aprobados. Los registros de despliegue se adjuntan en Evidencias. La automatización compila Java en CI; su ejecución real HTTP se registró mediante BodegaNorte.Tools en el entorno aislado.')
write(d,'12_Plan_Integracion_Continua.docx')

d=new('Anexo de validación final Bodega Norte','Resultados y capturas a color del Grupo 10')
text(d,STATUS);section(d,'Cómo leer las evidencias');text(d,'Las capturas siguientes provienen de recorridos reales en Chrome con contextos nuevos y datos ficticios locales. Muestran acciones ejecutadas durante las pruebas y conservan el aviso de demostración. Las operaciones PostgreSQL tienen registro propio de API y base aislada. Las ocho consultas IBM son una validación real separada. No se incluyen contraseñas ni claves privadas.')
grid(d,['Registro','Qué demuestra'],[('validacion_final_2026-10-11.json','17 casos HTTP y PostgreSQL; base principal sin cambios y limpieza de esquemas propios.'),('validacion_interfaz.json','11 recorridos y 14 capturas; cero errores y advertencias.'),('validacion_watson_2026-10-11.json','8 consultas con sus respuestas en la instancia IBM real.'),('despliegue_final_2026-10-11.json','Readiness HTTP 200; productos y estado del asistente HTTP 401 sin sesión; CORS permitido.')],[7,9])
text(d,'El respaldo demo y el CSV descargados durante las pruebas se adjuntan como artefactos de datos ficticios. La constancia de venta es interna y no representa un comprobante tributario.')
for i,shot in enumerate(ui['screenshots'],1):
    d.add_page_break();section(d,f'Evidencia {i} {shot["caption"]}')
    text(d,'Entorno: demostración local aislada. Fecha de registro: 11 de octubre de 2026. Fuente: recorrido automatizado de interfaz con interacción real.')
    b.image(d,str((EVID/shot['file']).relative_to(ROOT)),f'Figura {i} {shot["caption"]}',2.8 if 'movil' in shot['file'] else 6.1,color=True)
    if i==10:text(d,'La captura muestra la validación del archivo y la confirmación RESTAURAR antes de aplicar la copia. El recorrido posterior comprobó la recuperación y su auditoría.')
    elif i==4:text(d,'La venta de prueba con Yape generó una constancia por S/ 13.00. No se realizó un pago bancario.')
    elif i==13:text(d,'Viewport móvil de 390 por 844 píxeles. Se verificó que el menú abre el módulo de venta y no hay desbordamiento horizontal.')
    elif i==14:text(d,'Esta evidencia muestra ayuda local, no una respuesta de IBM. Las consultas IBM reales están registradas en su archivo independiente.')
    else:text(d,'Resultado del recorrido: aprobado. Consulte la matriz para relacionar esta vista con la operación y sus controles.')
write(d,'Anexo_Validacion_Final_Grupo10.docx')

d=new('Guía del paquete final Bodega Norte','Archivos de entrega y pasos para ejecutar el proyecto')
section(d,'Entrega del curso');b.bullets(d,['Informe_Final_Grupo10.docx contiene el proyecto por capítulos y los recursos solicitados.','Anexo_Validacion_Final_Grupo10.docx contiene resultados y 14 capturas actuales a color.','La hoja Anexo1A conserva integrantes y avances propuestos; no acredita horas ni fechas de asistencia.','Los manuales, cronograma, backlog, matriz, Watson y guion están en Documentación.','Código Visual Studio MVC, frontend, scripts SQL, prueba Java y acciones IBM se adjuntan completos sin credenciales privadas.'])
section(d,'Ejecución');steps(d,['Extraiga el ZIP completo. Abra una terminal en 06_Codigo/EFSR14908_Bodega_Norte.','Ejecute pnpm install --frozen-lockfile y pnpm run dev para la demo.','Para MVC ejecute pnpm run build:visualstudio, abra backend/BodegaNorte.sln y configure User Secrets siguiendo la guía.','Para revisión abra primero el informe, después el anexo y la matriz. Los JSON conservan resultados técnicos y los archivos de capturas son sus originales.'])
section(d,'Web publicada');text(d,'https://jaguar009.github.io/efsr14908-bodega-norte/');text(d,'Backend: https://bodega-norte-grupo10.onrender.com');text(d,'El ingreso exige una cuenta confirmada y habilitada. El proyecto académico está publicado; las pruebas no acreditan operación en una bodega real. IBM usa Trial gratuito de 30 días y conserva ayuda local independiente.')
write(d,'00_Guia_del_Paquete_Final.docx')
print(json.dumps({'documents':12,'captures':len(ui['screenshots']),'backendCases':len(backend['cases']),'ibmCases':len(ibm['cases'])}))
