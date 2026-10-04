from pathlib import Path
from tempfile import NamedTemporaryFile
from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageOps


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "word"
OUT.mkdir(parents=True, exist_ok=True)

COURSE = "Experiencias Formativas en Situaciones Reales de Trabajo III (EFSR14908)"
PROJECT = "Sistema de inventario y ventas Bodega Norte"
TEACHER = "Judith Jenny Jimenez Monago"
PERIOD = "Ciclo cuarto | Periodo 2026 | Semestre no consignado en los documentos recibidos"
BLUE = "000000"
BLUE2 = "000000"
LIGHT = "FFFFFF"
GRAY = "000000"
TEXT = "000000"
GROUP = "Grupo 10"
COORDINATOR = "Gonzalo Yoel Choque Becerra"
TEAM = [
    "Adrian Alexander Berrocal Villanueva",
    "Lalo Alessandro Eugenio Montalvo",
    "Juan Diego Romero Peralta",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def style_document(doc):
    for sec in doc.sections:
        sec.page_width = Cm(21)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(3)
        sec.bottom_margin = Cm(3)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.keep_together = True
    normal.paragraph_format.widow_control = True
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(11)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("000000")
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.0
        if name == "Title":
            ppr = st._element.get_or_add_pPr()
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)
    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(11)
        st.font.color.rgb = RGBColor.from_string("000000")
        st.paragraph_format.line_spacing = 1.0
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "BODEGA NORTE | EFSR14908"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.name = "Arial"
        header.runs[0].font.size = Pt(11)
        header.runs[0].font.color.rgb = RGBColor.from_string("000000")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("EFSR14908 | Página ")
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string("000000")
        add_page_field(footer)


def title(doc, text, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("000000")
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(subtitle)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string("000000")


def cover(doc, doc_title, deliverable=None):
    doc.add_paragraph("", style="Normal")
    title(doc, "BODEGA NORTE", doc_title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COURSE)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("000000")
    p = doc.add_paragraph(TEACHER)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(PERIOD)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(GROUP)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    p = doc.add_paragraph(f"Coordinador: {COORDINATOR}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para(doc, "Integrantes: " + "; ".join(TEAM))
    para(doc, "Fecha de elaboración: 4 de octubre de 2026")
    doc.add_page_break()


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.65)
        p.add_run(f"{index}. {item}")


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0]
    hdr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(value)
        set_cell_shading(cell, "000000")
        set_cell_borders(cell, "000000")
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
    for ridx, row in enumerate(rows):
        new_row = t.add_row()
        new_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = new_row.cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_shading(cells[i], "FFFFFF")
            set_cell_borders(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)
                    r.font.color.rgb = RGBColor.from_string(TEXT)
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Cm(width)
    doc.add_paragraph("")
    return t


def image(doc, path, caption=None, width=6.5):
    pth = ROOT / path
    if not pth.exists():
        para(doc, f"Evidencia pendiente de insertar: {path}")
        return
    with Image.open(pth) as source:
        bw = ImageOps.autocontrast(source.convert("L")).point(lambda value: 255 if value > 190 else 0, mode="1")
        with NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            temporary_image = Path(tmp.name)
        bw.save(temporary_image, format="PNG")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(temporary_image), width=Inches(width))
    temporary_image.unlink(missing_ok=True)
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        r.italic = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string("000000")


def save(doc, filename):
    style_document(doc)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string("000000")
    for table_obj in doc.tables:
        for row_index, row in enumerate(table_obj.rows):
            for cell in row.cells:
                header_cell = row_index == 0 and table_obj.columns[0].cells[0].text != "Curso"
                set_cell_shading(cell, "000000" if header_cell else "FFFFFF")
                set_cell_borders(cell, "000000")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
                        run.font.size = Pt(11)
                        run.font.color.rgb = RGBColor.from_string("FFFFFF" if header_cell else "000000")
    for section in doc.sections:
        for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
            for run in paragraph.runs:
                run.font.name = "Arial"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor.from_string("000000")
    doc.save(OUT / filename)



def doc_01():
    d = Document(); cover(d, "Informe inicial del proyecto", "Propuesta de sistema de inventario y ventas")
    h(d, "Introducción")
    para(d, "Este informe define el proyecto académico Bodega Norte, un prototipo web para organizar productos, existencias y ventas de una pequeña bodega representada como caso de estudio. Se describen el problema supuesto, el contexto SEPTE, los objetivos, el alcance, la arquitectura propuesta, el equipo y los beneficiarios. El escenario no corresponde a una empresa participante: no se adjuntaron entrevistas, mediciones de tiempos ni registros operativos reales.")
    para(d, "El planteamiento se presenta como propuesta académica y no certifica que sus actividades se hayan ejecutado en una fecha anterior. El estado técnico al cierre del paquete se documenta en el informe final y la matriz de pruebas.")

    h(d, "Capítulo I Diagnóstico del problema")
    h(d, "1.1 Diagnóstico situacional", 2)
    para(d, "Bodega Norte representa una bodega minorista ficticia de Lima Metropolitana. Para delimitar el ejercicio se supone que los datos de productos, ventas y existencias podrían mantenerse en apuntes o herramientas independientes. Esa condición se utiliza como problema de diseño y deberá contrastarse con una persona usuaria antes de aplicar la solución en un comercio real.")
    para(d, "La necesidad que se aborda es consultar el catálogo y el stock con una referencia común y registrar ventas de manera que la cantidad disponible se actualice según las unidades vendidas. El prototipo trabaja con productos y operaciones de demostración; sus datos no son una línea base del sector ni de un negocio real.")
    h(d, "1.2 Adversidades potenciales del escenario", 2)
    table(d, ["Situación potencial", "Consecuencia que se busca prevenir"], [
        ("Registros de productos y existencias separados.", "Duplicidad, desfase o dificultad para ubicar un dato vigente."),
        ("Venta y actualización de stock sin una regla común.", "Diferencias entre las unidades anotadas y las disponibles."),
        ("Falta de referencia de stock mínimo.", "Demora para detectar productos que requieren revisión o reposición."),
        ("Ventas sin consulta centralizada.", "Mayor esfuerzo al revisar operaciones recientes."),
        ("Dependencia de cálculos manuales.", "Errores de digitación o cálculo que deben prevenirse con validaciones."),
        ("Reportes preparados por separado.", "Información operativa menos accesible para el usuario del caso."),
    ])
    para(d, "Estas adversidades describen riesgos del escenario planteado, no hallazgos de una investigación de campo.")

    h(d, "1.3 Análisis SEPTE", 2)
    h(d, "Aspecto social", 3)
    para(d, "El sistema se diseña para un usuario de mostrador que necesita realizar consultas y registrar operaciones con pocos pasos. La facilidad de uso y la aceptación del personal no se han medido; deberán revisarse mediante una prueba guiada con usuarios del comercio si el caso se valida fuera del aula.")
    h(d, "Aspecto económico", 3)
    para(d, "Las Mipyme representaron el 99,3 % de las empresas formales del país en 2024 y las microempresas el 93,8 %. Esta composición aporta contexto para estudiar herramientas de gestión accesibles; no demuestra una necesidad particular de Bodega Norte, que es un caso simulado [1].")
    h(d, "Aspecto político y normativo", 3)
    para(d, "Si una versión futura tratara datos personales, deberá atender la Ley N.° 29733 y el reglamento vigente aprobado por el Decreto Supremo N.° 016-2024-JUS. La Autoridad Nacional de Protección de Datos Personales informa que dicho reglamento rige desde el 31 de marzo de 2025 [2]. El prototipo usa registros ficticios. Se implementaron autenticación, roles y auditoría; su operación remota sigue pendiente de validar.")
    h(d, "Aspecto tecnológico", 3)
    para(d, "El INEI publica indicadores sobre tecnologías de información en empresas, útiles para contextualizar la adopción digital sin extrapolar sus resultados a una bodega específica [3]. Por ello, el proyecto considera una demostración local y una arquitectura con backend MVC y Supabase PostgreSQL, cuya configuración se explica por separado.")
    h(d, "Aspecto ecológico", 3)
    para(d, "El proyecto puede ejecutarse en equipos existentes y no requiere hardware especializado para la demostración. Los equipos electrónicos que lleguen al final de su vida útil deben gestionarse conforme al régimen peruano de residuos de aparatos eléctricos y electrónicos [4]. No se atribuye al prototipo una reducción ambiental medida.")
    h(d, "1.4 Justificación", 2)
    para(d, "La propuesta concentra catálogo, existencias y ventas en una solución que permite mostrar cómo se relacionan esas operaciones. Desde el punto de vista técnico, separar interfaz, lógica del backend y persistencia facilita revisar y mantener los componentes. Desde el punto de vista formativo, el caso permite aplicar análisis, desarrollo, integración, pruebas y documentación sin usar datos de un negocio real.")
    para(d, "El beneficio esperado es facilitar la consulta y reducir inconsistencias dentro del flujo de demostración. No se promete una reducción porcentual del tiempo de atención, ahorro económico ni impacto comercial, porque no se cuenta con mediciones comparativas.")

    h(d, "Capítulo II Descripción del proyecto")
    h(d, "2.1 Descripción general y arquitectura", 2)
    para(d, "Bodega Norte es una aplicación web académica con interfaz React, backend ASP.NET Core 8 MVC y persistencia prevista en Supabase PostgreSQL. La vista de demostración también puede usar almacenamiento local del navegador para recorrer el flujo sin servidor. La demo local y el modo conectado a la API son alternativas distintas; sus datos no se sincronizan automáticamente.")
    table(d, ["Capa", "Responsabilidad"], [
        ("Interfaz React", "Consultar productos, revisar existencias y preparar ventas de demostración."),
        ("API ASP.NET Core", "Recibir solicitudes, validar reglas de negocio y devolver resultados."),
        ("Servicio y repositorio", "Organizar casos de uso y acceso a datos."),
        ("Supabase PostgreSQL", "Guardar categorías, proveedores, productos, ventas y líneas de venta mediante Npgsql."),
    ])
    h(d, "2.2 Objetivos", 2)
    h(d, "Objetivo general", 3)
    para(d, "Desarrollar y documentar un prototipo web que centralice el catálogo, permita registrar ventas y actualice las existencias de Bodega Norte en un entorno demostrativo local.")
    h(d, "Objetivos específicos e indicadores", 3)
    table(d, ["Objetivo", "Indicador de verificación"], [
        ("Gestionar al menos 10 productos de prueba con categoría, precio, stock y mínimo.", "Registros de prueba disponibles y campos consultables."),
        ("Registrar una venta válida y actualizar el stock de sus productos.", "Prueba válida: líneas guardadas y cantidades recalculadas."),
        ("Rechazar una venta que exceda las existencias.", "Prueba de exceso: validación y stock sin cambios."),
        ("Facilitar la consulta de stock bajo y exportar inventario.", "Filtro visible y CSV con los campos previstos."),
    ])
    para(d, "Los indicadores se verifican con datos de demostración. La matriz final distingue las pruebas técnicas realizadas de la aceptación manual pendiente.")

    h(d, "2.3 Alcance", 2)
    h(d, "Incluye", 3)
    bullets(d, ["Catálogo e inventario con búsqueda y referencia al stock mínimo.", "Punto de venta demostrativo con carrito, cantidades y medios de pago simulados.", "Actualización transaccional de stock y almacenamiento de venta y líneas en la API.", "Indicadores de demostración, exportación CSV, esquema SQL, datos de ejemplo, diagramas, código y manuales."])
    h(d, "Fuera del alcance", 3)
    bullets(d, ["Operación productiva, despliegue público o conexión a una empresa real.", "Multi-sucursal, compras, contabilidad, facturación electrónica, SUNAT y pagos bancarios.", "Copias automáticas y operación en una empresa real. El alcance actualizado incluye autenticación, roles, auditoría y exportación de respaldo JSON.", "Conexión productiva a IBM Watson Assistant; la ayuda disponible es una demo local."])

    h(d, "2.4 Propuesta de valor y sostenibilidad", 2)
    para(d, "El valor del proyecto está en adaptar un prototipo editable al caso académico y explicar sus componentes. No se han medido costos, rendimiento ni funciones de proveedores comerciales, por lo que no se afirma superioridad frente a plataformas existentes. No se plantea monetización ni se ha calculado un costo comercial.")
    para(d, "Para una posible continuidad se propone reutilizar equipo disponible, mantener datos ficticios en las demostraciones y definir autenticación, respaldos, seguridad y privacidad antes de cualquier uso real.")
    h(d, "2.5 Ubicación y entorno", 2)
    para(d, "El contexto comercial es una bodega ficticia de Lima Metropolitana; no se consigna dirección ni razón social real. El entorno técnico requiere navegador, Node.js para compilar la interfaz, .NET 8 para el backend y un proyecto Supabase PostgreSQL configurado para persistencia. La guía de instalación detalla los pasos.")
    h(d, "2.6 Organización del equipo", 2)
    para(d, "La publicación del Grupo 10 identifica a Gonzalo Yoel Choque Becerra como coordinador. Los roles siguientes proceden del Entregable 1 proporcionado y sirven como distribución de referencia; no certifican el historial individual de trabajo.")
    table(d, ["Rol declarado", "Integrante", "Responsabilidad de referencia"], [
        ("Líder del proyecto", COORDINATOR, "Coordinar el plan, acuerdos y consolidación del entregable."),
        ("Desarrollador frontend", TEAM[0], "Organizar interfaz, navegación y componentes visuales."),
        ("Responsable de base de datos", TEAM[1], "Revisar esquema, relaciones y persistencia SQL."),
        ("Desarrollador backend", TEAM[2], "Organizar API y operaciones de inventario y ventas."),
    ])
    h(d, "2.7 Beneficiarios", 2)
    table(d, ["Beneficiario", "Beneficio esperado, sujeto a validación"], [
        ("Personal de caja o administración", "Consultar productos y existencias y registrar ventas en la demo."),
        ("Clientes del comercio representado", "Podrían recibir atención más ordenada si la solución se valida en un negocio real."),
        ("Proveedores del comercio representado", "La referencia de existencias podría apoyar la planificación de reposición."),
        ("Grupo 10", "Aplicar y documentar competencias de desarrollo y gestión del proyecto."),
    ])
    h(d, "Bibliografía", 1)
    bullets(d, [
        "[1] Ministerio de la Producción. Las Mipyme en cifras 2024. OGEIEE, Lima, 2025. https://ogeiee.produce.gob.pe/index.php/en/shortcode/oee-documentos-publicaciones/publicaciones-anuales/item/download/2735_6390d933ea51b8a09fdd4a2dae545f5c",
        "[2] Autoridad Nacional de Protección de Datos Personales. Nuevo Reglamento de Protección de Datos Personales. https://www.gob.pe/institucion/anpd/campanas/128319-nuevo-reglamento-de-proteccion-de-datos-personales",
        "[3] Instituto Nacional de Estadística e Informática. Perú: Tecnologías de Información y Comunicación en las Empresas, 2023. https://www.gob.pe/institucion/inei/informes-publicaciones/7820458-peru-tecnologias-de-informacion-y-comunicacion-en-las-empresas-2023",
        "[4] Ministerio del Ambiente. Decreto Supremo N.° 009-2019-MINAM. https://www.gob.pe/institucion/minam/normas-legales/354138-009-2019-minam",
    ])
    save(d, "01_Entregable_1_Informe_Inicial.docx")


def doc_02():
    d = Document(); cover(d, "Informe de avance del proyecto", "Entregable de referencia para el hito de avance")
    h(d, "1. Estado verificable del prototipo")
    para(d, "Este documento describe el estado del código disponible al preparar el paquete, con fecha 4 de octubre de 2026. No certifica que el equipo haya presentado exactamente este avance en la semana 10 ni calcula un porcentaje de ejecución según la rúbrica.")
    table(d, ["Área", "Estado del código", "Límite de la evidencia"], [
        ("Interfaz", "Hay vistas para inicio, venta, productos, inventario, reportes y configuración.", "Pendiente aceptación manual del Grupo 10."),
        ("Ventas API", "El código valida stock y descuenta unidades dentro de una transacción PostgreSQL.", "Conexión y prueba remota pendientes de configurar."),
        ("Base de datos", "Esquema SQL con categorías, proveedores, productos, ventas y líneas.", "Datos de ejemplo; sin uso de una bodega real."),
        ("Reportes", "La interfaz calcula reportes por período y exporta inventario, ventas y movimientos a CSV.", "Descarga pendiente de prueba visual por el equipo."),
        ("Asistente", "La demo local ofrece respuestas de muestra; contrato Watson documentado.", "Servicio Watson no conectado."),
    ])
    h(d, "2. Historias priorizadas")
    table(d, ["Historia", "Criterio de aceptación", "Estado"], [
        ("Consultar productos y existencias.", "La vista y `GET /api/products` exponen código, categoría, stock, mínimo y precio.", "Implementado; revisión de interfaz pendiente."),
        ("Registrar una venta.", "El código API valida los datos, registra líneas y descuenta stock.", "Prueba anterior en SQL Server; pendiente validar la migración Supabase."),
        ("Rechazar exceso de venta.", "El código API devuelve HTTP 400 y no altera el stock.", "Prueba anterior en SQL Server; pendiente validar la migración Supabase."),
        ("Identificar stock bajo.", "Se considera bajo cuando stock es menor o igual al mínimo.", "Implementado en interfaz; revisión manual pendiente."),
        ("Exportar inventario.", "El CSV contiene código, nombre, categoría, stock, mínimo, precio y proveedor.", "Incluido en código; aceptación manual pendiente."),
        ("Consultar ayuda.", "El asistente local ofrece respuesta de muestra.", "Demo local; Watson pendiente."),
    ])
    h(d, "3. Arquitectura")
    para(d, "La aplicación React funciona como vista. Los controladores ASP.NET reciben operaciones; BodegaService valida los casos de uso y BodegaRepository ejecuta consultas y transacciones PostgreSQL con Npgsql. El modo demo utiliza localStorage y no necesita el backend. Ambos modos se documentan por separado.")
    image(d, "docs/diagramas/arquitectura_bw.png", "Figura 1. Flujo de componentes", 6.1)
    h(d, "4. Responsabilidades del equipo")
    table(d, ["Rol declarado en el Entregable 1", "Integrante"], [
        ("Líder del proyecto", COORDINATOR),
        ("Desarrollador frontend", TEAM[0]),
        ("Responsable de base de datos", TEAM[1]),
        ("Desarrollador backend", TEAM[2]),
    ])
    para(d, "Los roles provienen del Entregable 1 entregado por el estudiante. Este informe no añade un historial de tareas individuales, actas de reunión ni firmas.")
    h(d, "5. Validaciones ejecutadas")
    table(d, ["Validación", "Resultado"], [
        ("Frontend: npm run build", "Correcto el 28-09-2026; 16 módulos transformados."),
        ("Backend: dotnet build", "Correcto el 28-09-2026; cero advertencias y cero errores."),
        ("Venta y descuento de stock", "Prueba histórica SQL Server; no valida PostgreSQL."),
        ("Venta por encima del stock", "Prueba histórica SQL Server; no valida PostgreSQL."),
        ("Interfaz, responsive, CSV y aceptación de usuario", "Pendiente de revisión del Grupo 10."),
    ])
    para(d, "El registro detallado se encuentra en `docs/evidencias/validacion_tecnica_2026-09-28.md`. La prueba HTTP documentada corresponde a la versión histórica SQL Server. El backend Npgsql actual ya compila, pero su conexión Supabase aún requiere la cadena de conexión del proyecto nuevo.")
    h(d, "6. Plan de integración continua")
    para(d, "El workflow `.github/workflows/ci.yml` compila los dos modos del frontend, compila la solución .NET y verifica que el programa Java compile. El workflow de GitHub Pages publica solo la demo estática. La prueba HTTP de venta con PostgreSQL queda pendiente de ejecutar contra el proyecto Supabase de desarrollo.")
    h(d, "7. Acciones antes de una entrega académica")
    bullets(d, ["Revisar los flujos de interfaz y guardar evidencia con ejecutor y fecha reales.", "Confirmar el calendario final del curso; el sílabo 2026 y el Anexo 4 Nivel 2 indican semanas distintas.", "Validar los roles y los criterios de aceptación con el equipo.", "Mantener sin datos personales reales la base de demostración y completar seguridad antes de producción."])
    h(d, "Actualización operativa al 04-10-2026")
    para(d, "Se implementaron Supabase Auth y permisos administrador/cajero, ajustes con motivo y control de versión, venta con identificador de reintento, anulaciones con devolución única de stock, historial y constancia interna, proveedores y configuración persistentes, auditoría y respaldo JSON. Los reportes usan fecha de Lima y detalles de venta; no se inventan costos históricos.")
    para(d, "Validación técnica: 18 pruebas locales aprobadas; frontend y .NET compilan, con cero errores y advertencias en .NET; Java compila. La revisión de navegador se realizó en modo demo. La conexión privada del backend, integración HTTP PostgreSQL, despliegue Render y aceptación del Grupo 10 están pendientes. La evidencia detallada está en docs/evidencias/validacion_tecnica_2026-10-04.md.")
    h(d, "8. Evidencias visuales")
    image(d, "docs/evidencias/EVID_visual_escritorio.png", "Evidencia visual de referencia en escritorio", 6.1)
    image(d, "docs/evidencias/EVID_visual_movil.png", "Evidencia visual de referencia en móvil", 3.2)

    save(d, "02_Entregable_2_Avance_50.docx")


def doc_03():
    d = Document(); cover(d, "Informe final del proyecto", "Bodega Norte y gestión de inventario y ventas")
    h(d, "1. Resumen")
    para(d, "Bodega Norte es un prototipo web académico para organizar productos, existencias y ventas de una bodega simulada. El paquete incluye una interfaz React, una alternativa de demostración con almacenamiento local y un backend ASP.NET Core MVC preparado para Supabase PostgreSQL. El alcance verificable comprende el catálogo, el punto de venta, el control de stock, los reportes CSV, el esquema de datos, los diagramas, la guía de instalación, el manual de usuario y los artefactos de gestión. No se presenta como una solución desplegada en una bodega real. La validación funcional del equipo y de un usuario de negocio queda registrada por separado de la compilación y de las pruebas técnicas ejecutadas.")
    h(d, "2. Introducción")
    para(d, "El caso de estudio representa a una pequeña bodega que necesita consultar productos, revisar existencias y registrar ventas sin duplicar cálculos en cuadernos o archivos independientes. Bodega Norte propone concentrar esas operaciones en una aplicación web. La organización y la operación comercial del caso son simuladas; no se realizaron entrevistas ni mediciones en una empresa real, por lo que los resultados describen el prototipo y no un impacto productivo observado.")
    para(d, "El informe sigue la estructura solicitada para EFSRT III. Describe el diagnóstico SEPTE, los objetivos, la propuesta de valor, el alcance, la gestión del proyecto, el producto, las pruebas y los recursos entregados. Las fechas de evaluación aparecen como semanas relativas porque los documentos del curso tienen referencias que no coinciden y el calendario final debe confirmarse en el aula virtual.")

    h(d, "3. Diagnóstico del problema")
    h(d, "3.1 Contexto del caso", 2)
    para(d, "Bodega Norte es una organización ficticia ubicada en Lima Metropolitana para fines académicos. El problema de trabajo asumido es la posible dispersión de los registros de productos, ventas y stock. El equipo no recibió datos de una bodega real ni resultados de entrevistas; por ello, esta descripción constituye el escenario del prototipo y debe contrastarse con usuarios antes de una implementación productiva.")
    h(d, "3.2 Análisis SEPTE", 2)
    h(d, "Factor social", 3)
    para(d, "La atención en mostrador depende de que el personal pueda verificar productos y registrar operaciones con rapidez. En este proyecto se trata como una necesidad del escenario, no como una conclusión estadística ni como un resultado de entrevistas. La aceptación social del sistema deberá evaluarse con el encargado y con usuarios reales.")
    h(d, "Factor económico", 3)
    para(d, "El Ministerio de la Producción reporta que las Mipyme representaron 99,3 % de las empresas formales del país en 2024; las microempresas fueron 2 201 422, equivalentes a 93,8 % del total formal. Estas cifras describen el contexto empresarial nacional y respaldan la pertinencia de estudiar herramientas de bajo costo para procesos de pequeños comercios, pero no prueban por sí mismas la situación de una bodega específica. [[FN1]]")
    h(d, "Factor político y normativo", 3)
    para(d, "Si una aplicación de ventas llegara a tratar datos personales, el responsable tendría que cumplir la Ley N.° 29733 y su reglamento vigente, aprobado por el Decreto Supremo N.° 016-2024-JUS. La Autoridad Nacional de Protección de Datos Personales informó que el reglamento entró en vigor el 31 de marzo de 2025. El prototipo usa datos ficticios, incorpora autenticación y no mantiene un registro de clientes reales; antes de producción se deben definir finalidad, acceso, conservación y medidas de seguridad. [[FN2]]")
    h(d, "Factor tecnológico", 3)
    para(d, "El INEI publica indicadores de TIC en empresas desagregados por tamaño y actividad económica. Esta fuente permite contextualizar la adopción tecnológica, pero el proyecto no extrapola sus resultados a todas las bodegas ni presupone conexión continua a Internet. Por esa razón se mantiene una demo local y una ruta con backend MVC y Supabase PostgreSQL, que requieren configuraciones distintas. [[FN3]]")
    h(d, "Factor ecológico", 3)
    para(d, "El uso de computadoras y otros equipos puede generar residuos de aparatos eléctricos y electrónicos al final de su vida útil. El régimen especial peruano establece responsabilidades para su gestión y disposición. Para el caso de estudio se recomienda reutilizar equipos disponibles y entregar los equipos descartados a un sistema autorizado; el uso de registros digitales no se presenta como una reducción ambiental medida. [[FN4]]")

    h(d, "4. Objetivos")
    h(d, "4.1 Objetivo general", 2)
    para(d, "Desarrollar y documentar un prototipo web que centralice el catálogo, permita registrar ventas y actualice las existencias de Bodega Norte en un entorno demostrativo local.")
    h(d, "4.2 Objetivos SMART", 2)
    table(d, ["Objetivo", "Medida y meta", "Plazo y estado"], [
        ("Implementar el catálogo del prototipo.", "Gestionar al menos 10 productos de demostración con categoría, precio, stock y mínimo.", "Al cierre del proyecto. Funcionalidad incluida; la aceptación de interfaz queda pendiente de registro del equipo."),
        ("Registrar ventas sin generar stock negativo.", "La venta válida debe guardar sus líneas y descontar las cantidades en una transacción; una cantidad superior al stock debe rechazarse.", "Al cierre del proyecto. Regla implementada en el backend; resultado de verificación indicado en la matriz."),
        ("Facilitar consultas operativas.", "Mostrar productos bajo el mínimo y permitir exportar el inventario a CSV.", "Al cierre del proyecto. Filtro y exportación incluidos; la validación visual del equipo queda pendiente."),
    ])
    para(d, "Las metas se miden sobre datos de demostración. No se fija una reducción porcentual del tiempo de atención porque no existe una línea base observada en una bodega real.")

    h(d, "5. Justificación del proyecto")
    para(d, "El prototipo permite mostrar en una misma solución el catálogo, las existencias y la venta de mostrador. En el modo con base de datos, el backend valida los datos y procesa la venta dentro de una transacción; en el modo de demostración, el navegador conserva datos de ejemplo para recorrer la interfaz. Esta separación ayuda a explicar el alcance técnico y evita presentar la demo local como un servicio productivo.")
    h(d, "5.1 Beneficiarios", 2)
    table(d, ["Tipo", "Beneficiario", "Beneficio esperado"], [
        ("Directo", "Personal de caja o administrador de la bodega simulada.", "Consultar catálogo y existencias, registrar una venta y revisar alertas."),
        ("Directo", "Equipo del Grupo 10.", "Aplicar análisis, diseño, programación, gestión y documentación del curso."),
        ("Indirecto", "Clientes del comercio representado.", "Podrían recibir atención más ordenada si el producto se validara en una operación real."),
        ("Indirecto", "Proveedores del comercio representado.", "La información de stock podría apoyar la planificación de reposición, sujeto a uso y validación."),
    ])
    h(d, "5.2 Modelo Canvas", 2)
    table(d, ["Bloque", "Propuesta para el caso de estudio"], [
        ("Segmentos", "Bodegas pequeñas y personal que registra ventas o administra productos."),
        ("Propuesta de valor", "Catálogo, venta e inventario en una interfaz sencilla con alertas y reportes básicos."),
        ("Canales", "Navegador en equipo local; posible despliegue futuro sujeto a configuración y seguridad."),
        ("Relación con usuarios", "Uso guiado por módulos, manual y ayuda local de demostración."),
        ("Actividades clave", "Mantener catálogo, registrar ventas, revisar stock y respaldar información."),
        ("Recursos clave", "Aplicación web, backend, base de datos, equipo de cómputo y documentación."),
        ("Aliados clave", "Equipo del curso, docente-monitor y, si se valida el caso, responsables del comercio."),
        ("Costos", "Desarrollo, equipo existente, conectividad y mantenimiento; no se estimó un costo comercial."),
        ("Ingresos o beneficios", "No se plantea monetización. Se espera facilitar el control operativo, sujeto a medición futura."),
    ])

    h(d, "6. Definición y alcance del proyecto")
    h(d, "6.1 Funcionamiento", 2)
    numbered(d, [
        "El usuario consulta el catálogo y selecciona productos para una venta.",
        "La interfaz calcula cantidades y total y envía la solicitud al backend MVC conectado a Supabase PostgreSQL.",
        "El servicio valida el método de pago, las cantidades y la disponibilidad de cada producto.",
        "El repositorio registra la venta y sus líneas y descuenta el stock dentro de una transacción serializable.",
        "La interfaz vuelve a consultar los datos y muestra el resultado. En la demo sin servidor, el estado se guarda en localStorage.",
    ])
    h(d, "6.2 Incluye", 2)
    bullets(d, [
        "Panel de inicio, catálogo de productos, inventario con búsqueda y filtro de stock bajo.",
        "Punto de venta con carrito, cantidades, total y medios de pago de demostración.",
        "Proveedores persistentes con contactos, vínculos y restricciones. Compras no forma parte del alcance.",
        "Reportes operativos y exportación CSV del inventario.",
        "Backend ASP.NET Core 8 MVC, repositorio Npgsql, esquema PostgreSQL y datos de demostración.",
        "Diagramas, guía de instalación, manual de usuario, manual técnico, plan de trabajo y matriz de pruebas.",
    ])
    h(d, "6.3 Fuera del alcance", 2)
    bullets(d, [
        "Operación productiva con una empresa real, multi-sucursal, inventario de compras o contabilidad.",
        "Facturación electrónica, SUNAT, pasarelas de pago, conexión bancaria o comprobantes fiscales.",
        "Copias automatizadas y despliegue operativo aún no verificado. Autenticación, roles y auditoría sí se implementaron.",
        "Conexión productiva con IBM Watson Assistant. Se entrega la especificación; la interfaz actual mantiene respuestas locales de demostración.",
    ])
    h(d, "6.4 Interesados", 2)
    table(d, ["Interesado", "Interés o participación"], [
        ("Docente-monitor", "Define y evalúa los entregables del curso."),
        ("Grupo 10", "Planifica, desarrolla, revisa y sustenta el prototipo."),
        ("Usuario de caja o administrador", "Perfil de usuario supuesto para diseñar los flujos; no se ha confirmado un participante real."),
        ("Clientes y proveedores", "Beneficiarios indirectos hipotéticos en la organización representada."),
    ])
    h(d, "6.5 Riesgos", 2)
    table(d, ["Riesgo", "Probabilidad e impacto", "Respuesta"], [
        ("Requisitos incompletos del caso ficticio", "Media / Media", "Validar flujos con un responsable real antes de extender el producto."),
        ("Pérdida o inconsistencia de datos locales", "Media / Alta", "Realizar respaldo de la base y no usar la demo local como única fuente de datos."),
        ("Acceso no autorizado en un despliegue", "Media / Alta", "Verificar las cuentas y roles implementados, HTTPS y auditoría en el despliegue antes de operar."),
        ("Fechas del curso no coincidentes entre documentos", "Alta / Alta", "Tomar la fecha publicada en el aula virtual y confirmar la semana final con la docente."),
        ("Integración Watson sin credenciales o servicio configurado", "Alta / Media", "Mantener el contrato documentado y etiquetar la conexión como pendiente."),
    ])
    h(d, "6.6 Viabilidad", 2)
    para(d, "Técnica: viable como prototipo con React, ASP.NET Core MVC y PostgreSQL en Supabase; la conexión remota del proyecto requiere configuración y validación. Operativa: los flujos son apropiados para demostración, pero no han sido aceptados por usuarios de una bodega real. Económica: no se calculó un costo de implementación ni se propone venta; el desarrollo usa herramientas y datos de demostración. Académica: el producto puede sustentarse con código, esquema, diagramas, manuales y una matriz de pruebas que distingue ejecuciones registradas de validaciones pendientes.")
    h(d, "6.7 Documentación entregada", 2)
    table(d, ["Documento o recurso", "Contenido"], [
        ("Informe final", "Diagnóstico, objetivos, Canvas, alcance, gestión, producto y anexos."),
        ("Cronograma y backlog", "Plan semanal, hitos, historias de usuario, responsables propuestos y revisión de fechas."),
        ("Guía de instalación", "Requisitos, ejecución demo, configuración del proyecto Supabase y User Secrets."),
        ("Manual de usuario", "Flujos de consulta, productos, venta, inventario, reportes y asistente local."),
        ("Manual técnico", "Arquitectura, entidades, contratos, reglas y limitaciones."),
        ("Matriz de pruebas", "Casos funcionales y técnicos, estado de ejecución y evidencias disponibles."),
        ("Código y base de datos", "Frontend, backend, esquema SQL y registros de demostración."),
        ("Diagramas y CI", "Arquitectura, modelo de datos y flujo propuesto de integración continua."),
    ])

    h(d, "7. Productos y entregables")
    table(d, ["Entregable", "Archivos"], [
        ("Aplicación web", "src/main.jsx, src/styles.css, src/api.js y public/."),
        ("Backend y solución Visual Studio", "backend/BodegaNorte.sln y backend/BodegaNorte.Api/."),
        ("Persistencia", "database/schema.sql, database/upgrade_operations.sql, database/seed.sql y database/README.md."),
        ("Diagramas", "docs/diagramas/arquitectura.drawio, arquitectura.svg, modelo_datos.svg y versiones monocromáticas."),
        ("Plan y pruebas", "docs/cronograma.md, docs/gestion/, docs/matriz_pruebas.md e integration/."),
        ("Recursos editables", "docs/word/ con informe, guía, manuales y materiales de sustentación."),
    ])

    h(d, "8. Planeamiento y control del proyecto")
    h(d, "8.1 Organización del equipo", 2)
    para(d, "Gonzalo Yoel Choque Becerra figura como coordinador en la publicación del Grupo 10. El Entregable 1 del equipo asigna los roles de líder del proyecto, desarrollador frontend, responsable de base de datos y desarrollador backend. Esa distribución se usa para organizar el plan y la sustentación; no constituye un registro de tareas individuales realizadas.")
    table(d, ["Rol declarado", "Integrante", "Alcance de la función"], [
        ("Líder del proyecto", COORDINATOR, "Seguimiento del cronograma, acuerdos y consolidación de la entrega."),
        ("Desarrollador frontend", TEAM[0], "Interfaz, navegación, componentes visuales y criterios de uso."),
        ("Responsable de base de datos", TEAM[1], "Esquema SQL, relaciones, consultas y persistencia."),
        ("Desarrollador backend", TEAM[2], "Controladores, servicios y operaciones de venta e inventario."),
    ])
    h(d, "8.2 Método de trabajo", 2)
    para(d, "El plan usa Scrum de forma académica: una lista priorizada de historias, incrementos breves, revisión de criterios de aceptación y registro de incidencias. Los archivos incluidos describen el proceso propuesto; no se adjuntaron actas de ceremonias ni registros históricos de cada sprint, por lo que el informe no afirma que esas reuniones se hayan realizado.")
    table(d, ["Historia", "Criterio de aceptación", "Estado en el prototipo"], [
        ("Como administrador, quiero consultar productos con stock y mínimo.", "El catálogo muestra los campos y permite buscar productos.", "Implementado en frontend y API; falta firma de aceptación del usuario."),
        ("Como cajero, quiero registrar una venta con método de pago.", "La venta guarda las líneas, calcula el total y descuenta stock.", "Implementado en modo SQL; ver resultados técnicos en la matriz."),
        ("Como responsable de inventario, quiero identificar stock bajo.", "Se consideran bajos los productos con stock menor o igual al mínimo.", "Implementado en la interfaz de demostración."),
        ("Como administrador, quiero exportar el inventario.", "Se descarga un CSV que puede abrirse en una hoja de cálculo.", "Incluido en Reportes; pendiente de aceptación manual del equipo."),
        ("Como usuario, quiero consultar la ayuda de la bodega.", "La interfaz ofrece respuestas para preguntas frecuentes.", "Demo local; Watson no está conectado."),
    ])
    h(d, "8.3 Plan semanal y Gantt resumido", 2)
    table(d, ["Semana relativa", "Tareas y subtareas", "Hito o evidencia"], [
        ("1", "Formar equipo, elegir problema y definir interesados.", "Grupo 10 y escenario Bodega Norte."),
        ("2–3", "Investigar contexto, plantear objetivos, justificación y alcance; preparar el primer informe.", "Entregable 1; el aula virtual indicó 21–27 de septiembre de 2026."),
        ("4", "Definir historias, flujo, criterios de aceptación y modelo de datos.", "Backlog y diagramas."),
        ("5–6", "Construir interfaz, catálogo, filtros y persistencia demo; preparar la base SQL.", "Producto base y esquema."),
        ("7–8", "Implementar venta, validación de stock, reportes y gestión de proveedores.", "Flujos de venta e inventario."),
        ("9", "Integrar módulos, corregir consistencia de datos y ejecutar pruebas de build/API.", "Registro de validación técnica."),
        ("10", "Presentar avance mínimo del 50 % si se confirma el cronograma del Anexo 4.", "Hito de avance y sustentación grupal."),
        ("11–12", "Completar manuales, matriz, anexos, revisión responsive y ensayo de defensa.", "Recursos editables y guion."),
        ("13", "Cerrar componentes, revisar entregables y sustentar si la semana permanece vigente.", "Entrega final según Anexo 4 Nivel 2; confirmar el día en el aula."),
    ])
    para(d, "Conflicto de calendario: el Anexo 4 Nivel 2 aportado indica avance en semana 10 y entrega final en semana 13; el sílabo 2026 aportado ubica la evaluación final en semana 7. Se conserva el plan de 13 semanas como referencia de implementación y no se fija una fecha de entrega final hasta confirmar el cronograma vigente de la docente.")

    h(d, "9. Producto")
    h(d, "9.1 Diagramas de análisis y diseño", 2)
    para(d, "El paquete adjunta el diagrama de arquitectura y el modelo entidad-relación en formatos editables y SVG. La arquitectura recorre la vista React, los controladores, el servicio de negocio, el repositorio Npgsql y Supabase PostgreSQL. El modelo de datos relaciona categorías y proveedores con productos, y ventas con sus líneas de detalle.")
    image(d, "docs/diagramas/arquitectura_bw.png", "Figura 1. Arquitectura del prototipo", 6.1)
    image(d, "docs/diagramas/modelo_datos_bw.png", "Figura 2. Modelo de datos SQL", 6.1)
    table(d, ["Flujo de arquitectura", "Entidades SQL actuales"], [
        ("React View → ASP.NET Controllers → BodegaService → BodegaRepository (Npgsql) → Supabase PostgreSQL", "categories, suppliers, products, sales, sale_items, members, settings, stock_movements y audit_log."),
        ("La respuesta vuelve del backend a la vista; la demo puede usar localStorage.", "Las claves foráneas vinculan productos con categorías/proveedores y líneas con productos/ventas."),
    ])
    h(d, "9.2 Módulos, modelos y componentes", 2)
    table(d, ["Componente", "Implementación actual"], [
        ("Panel de inicio", "Indicadores, ventas recientes, existencias y alertas de stock."),
        ("Productos e inventario", "Alta, edición, búsqueda, filtro y referencia al stock mínimo."),
        ("Punto de venta", "Carrito, cantidades, cálculo de total, método de pago y confirmación."),
        ("Reportes", "Indicadores por fecha de Lima y detalle registrado, excluyendo anulaciones; CSV de ventas e inventario."),
        ("Proveedores y configuración", "CRUD persistente de proveedores y configuración. Compras fuera del alcance."),
        ("Asistente", "Respuestas offline en interfaz. Watson Assistant se describe como integración pendiente."),
        ("Backend", "HomeController, ProductsController, SalesController, HealthController, BodegaService y BodegaRepository."),
    ])
    h(d, "9.3 Archivos de creación y configuración", 2)
    bullets(d, [
        "package.json, pnpm-lock.yaml y vite.config.js para construir la interfaz con npm.",
        "backend/BodegaNorte.sln, BodegaNorte.Api.csproj, Program.cs y appsettings.Development.json para ASP.NET Core.",
        "database/schema.sql, database/upgrade_operations.sql y database/seed.sql para crear entidades y cargar datos ficticios.",
        ".env.example con variables demostrativas; no contiene credenciales reales.",
        ".github/workflows/ para compilación y publicación de la demo estática.",
    ])
    h(d, "9.4 Validación y estado de pruebas", 2)
    para(d, "La matriz separa compilación, prueba de integración API y casos que requieren revisión manual de la interfaz. Solo se marcarán como ejecutados los casos con salida o evidencia guardada; las capturas visuales de referencia no se consideran por sí solas una prueba funcional.")
    table(d, ["Validación", "Resultado", "Evidencia"], [
        ("Compilación frontend", "Aprobada", "Registro de npm run build."),
        ("Compilación backend", "Aprobada", "Registro de dotnet build."),
        ("Venta API y descuento de stock", "Pendiente para Supabase PostgreSQL", "La prueba aprobada del paquete anterior fue sobre SQL Server y no valida esta migración."),
        ("Rechazo de venta por stock insuficiente", "Pendiente para Supabase PostgreSQL", "La prueba histórica SQL Server no valida esta migración."),
        ("Interfaz, reportes CSV y responsive", "Pendiente de aceptación manual del equipo", "Capturas en docs/evidencias/ como referencia visual."),
    ])

    h(d, "10. Recursos")
    bullets(d, [
        "Guía de instalación para modo demo y configuración de Supabase con User Secrets.",
        "Manual de usuario con los pasos para venta, inventario, productos y reportes.",
        "Manual técnico con arquitectura, modelo de datos, reglas, endpoints y limitaciones.",
        "Matriz de pruebas con criterios, estado y espacio para anotar ejecutor y fecha reales.",
        "Cronograma, backlog, plan de integración continua, diagramas editables y guion para la sustentación.",
    ])
    para(d, "El despliegue de producción requiere credenciales y configuración que no están incluidas. No se incorporan datos de estudiantes, clientes reales, claves ni firmas ficticias.")

    h(d, "11. Conclusiones")
    numbered(d, [
        "El prototipo organiza los flujos principales de productos, inventario y ventas en una interfaz y un backend local; los datos incluidos son demostrativos.",
        "El registro transaccional y las reglas de stock cubren el flujo técnico principal, mientras que la aceptación de usuario y la operación con una bodega real todavía requieren evidencia.",
        "El paquete reúne código, SQL, diagramas, manuales y documentación de gestión para que el Grupo 10 pueda revisar, ejecutar y sustentar el alcance con sus limitaciones explícitas.",
    ])
    h(d, "12. Recomendaciones")
    numbered(d, [
        "Confirmar con la docente la fecha final aplicable, debido a la diferencia entre el sílabo 2026 y el Anexo 4 Nivel 2.",
        "Validar el flujo con un usuario de una bodega real y registrar ejecutor, fecha, datos de prueba y resultados antes de afirmar impacto operativo.",
        "Antes de operar, validar autenticación, roles, auditoría y recuperación del respaldo en el entorno publicado; proteger credenciales y gestionar los datos personales.",
    ])
    h(d, "13. Glosario")
    table(d, ["Término", "Definición"], [
        ("API", "Interfaz que permite que la aplicación solicite operaciones al backend."),
        ("Backend", "Parte del sistema que procesa reglas y accede a la base de datos."),
        ("CSV", "Archivo de texto con campos separados, utilizable en hojas de cálculo."),
        ("MVP", "Versión mínima de un producto destinada a demostrar funciones prioritarias."),
        ("MVC", "Patrón que organiza responsabilidades en modelo, vista y controlador."),
        ("Scrum", "Marco de trabajo ágil que organiza el desarrollo en incrementos y revisiones."),
        ("Stock mínimo", "Umbral configurado para señalar una posible necesidad de reposición."),
        ("Transacción", "Unidad de operaciones de base de datos que se confirma en conjunto o se revierte."),
        ("Watson Assistant", "Servicio conversacional de IBM considerado para una integración futura."),
    ])
    h(d, "14. Bibliografía")
    bullets(d, [
        "Ministerio de la Producción. Las Mipyme en cifras 2024. OGEIEE, Lima, 2025. https://ogeiee.produce.gob.pe/index.php/en/shortcode/oee-documentos-publicaciones/publicaciones-anuales/item/download/2735_6390d933ea51b8a09fdd4a2dae545f5c",
        "Instituto Nacional de Estadística e Informática. Perú: Tecnologías de Información y Comunicación en las Empresas, 2023. https://www.gob.pe/institucion/inei/informes-publicaciones/7820458-peru-tecnologias-de-informacion-y-comunicacion-en-las-empresas-2023",
        "Autoridad Nacional de Protección de Datos Personales. Nuevo Reglamento de Protección de Datos Personales. https://www.gob.pe/institucion/anpd/campanas/128319-nuevo-reglamento-de-proteccion-de-datos-personales",
        "Ministerio del Ambiente. Decreto Supremo N.° 009-2019-MINAM. https://www.gob.pe/institucion/minam/normas-legales/354138-009-2019-minam",
        "Cibertec. Anexo 4, Informe de Proyecto ETI, Plan Nivel 2 para EFSRT III, IV y V; Plan de Implementación de EFSRT–ETI 2024; Sílabo EFSRT III, código 14908, periodo 2026; anuncio y consigna de la actividad en el aula virtual. Material entregado por el estudiante.",
        "Documentación del equipo. Código fuente, scripts SQL, diagramas, manuales y matriz de pruebas incluidos en este paquete.",
    ])
    h(d, "15. Anexos")
    h(d, "Anexo A. Integrantes del Grupo 10", 2)
    table(d, ["Integrante", "Participación declarada en la publicación del grupo"], [
        (COORDINATOR, "Coordinador"), (TEAM[0], "Integrante"), (TEAM[1], "Integrante"), (TEAM[2], "Integrante"),
    ])
    h(d, "Anexo B. Evidencia y componentes", 2)
    bullets(d, [
        "Capturas de referencia: docs/evidencias/EVID_visual_escritorio.png y EVID_visual_movil.png. Se adjuntan fuera de este informe para conservar el documento en blanco y negro.",
        "Diagramas editables y fuentes: docs/diagramas/arquitectura.drawio, arquitectura.svg, modelo_datos.svg y versiones blanco y negro en PNG y SVG.",
        "Código, base de datos, pruebas, manuales e informes: carpetas src/, backend/, database/, integration/ y docs/.",
    ])
    h(d, "Anexo C. Datos pendientes de ratificación por el equipo", 2)
    bullets(d, [
        "Confirmar el semestre específico en el calendario de 2026 y la fecha final en el aula virtual.",
        "Validar las responsabilidades propuestas y anotar aportes reales si la docente los solicita.",
        "Completar las actas o firmas de aceptación solo con la participación efectiva de las personas correspondientes.",
    ])

    h(d, "Actualización operativa al 04-10-2026")
    para(d, "Se implementaron Supabase Auth y permisos administrador/cajero, ajustes con motivo y control de versión, venta con identificador de reintento, anulaciones con devolución única de stock, historial y constancia interna, proveedores y configuración persistentes, auditoría y respaldo JSON. Los reportes usan fecha de Lima y detalles de venta; no se inventan costos históricos.")
    para(d, "Validación técnica: 18 pruebas locales aprobadas; frontend y .NET compilan, con cero errores y advertencias en .NET; Java compila. La revisión de navegador se realizó en modo demo. La conexión privada del backend, integración HTTP PostgreSQL, despliegue Render y aceptación del Grupo 10 están pendientes. La evidencia detallada está en docs/evidencias/validacion_tecnica_2026-10-04.md.")
    save(d, "03_Entrega_Final_Informe.docx")


def doc_04():
    d = Document(); cover(d, "Cronograma de entregas y evidencias", "Documento de control del proyecto")
    h(d, "1. Propósito")
    para(d, "Este plan organiza las tareas del prototipo por semanas relativas, con subtareas y productos verificables. No representa un registro histórico de asistencia ni acredita que cada actividad se haya ejecutado en la semana indicada.")
    h(d, "2. Hitos del curso")
    table(d, ["Hito", "Referencia recibida", "Contenido"], [
        ("Entregable 1", "Aula virtual: 21–27 de septiembre de 2026", "Informe del proyecto y hoja de seguimiento del grupo."),
        ("Avance del proyecto", "Semana 10 en Anexo 4 Nivel 2", "Avance mínimo de 50 % y sustentación grupal."),
        ("Evaluación final", "Semana 07 en Sílabo EFSRT III 2026", "Evaluación final del curso."),
        ("Entrega final", "Semana 13 en Anexo 4 Nivel 2", "Proyecto concluido, informe, componentes, recursos y sustentación."),
    ])
    h(d, "3. Gantt resumido")
    table(d, ["Semana", "Tareas y subtareas", "Hito / evidencia"], [
        ("1", "Formar el grupo; seleccionar el problema; identificar los perfiles de usuario.", "Grupo y tema acordados."),
        ("2–3", "Buscar fuentes; elaborar SEPTE; redactar objetivos SMART; justificar la propuesta y delimitar alcance.", "Entregable 1 en Word."),
        ("4", "Priorizar historias; definir criterios; diseñar arquitectura y entidades SQL.", "Backlog y diagramas."),
        ("5–6", "Construir interfaz base; catálogo; filtros; datos demo; esquema y carga SQL.", "Prototipo base ejecutable."),
        ("7–8", "Completar carrito y venta; validar cantidades; descontar stock; implementar reportes.", "Flujo venta e inventario."),
        ("9", "Integrar módulos; corregir datos de demostración; compilar frontend y backend; ejecutar API de integración.", "Registro de compilación e integración."),
        ("10", "Consolidar presentación de avance; explicar funcionalidades disponibles y pendientes.", "Hito de avance según Anexo 4."),
        ("11–12", "Preparar manuales; revisar accesibilidad y responsive; completar matriz y anexos; ensayar sustentación.", "Recursos y guion editables."),
        ("13", "Revisar el paquete completo; validar la entrega en aula virtual; realizar sustentación si corresponde.", "Hito final según Anexo 4."),
    ])
    h(d, "4. Control del trabajo")
    bullets(d, ["Mantener historias y tareas en una lista priorizada con criterio de aceptación.", "Registrar bloqueos, cambios y acuerdos con fecha y responsable reales.", "Conservar salidas de compilación y pruebas en docs/evidencias/.", "Revisar los entregables con el docente-monitor antes de la fecha vigente del aula virtual."])
    h(d, "5. Diferencias de calendario")
    para(d, "El anuncio del Entregable 1 muestra un plazo del 21 al 27 de septiembre de 2026. El Anexo 4 Nivel 2 indica avance en semana 10 y cierre en semana 13, mientras el sílabo 2026 ubica la evaluación final en semana 7. El plan conserva ambos hitos relativos para no descartar un requisito; el aula virtual y la docente deben confirmar la fecha que rige la entrega final.")
    save(d, "04_Cronograma.docx")


def doc_05():
    d = Document(); cover(d, "Manual de usuario", "Guía de operación del sistema Bodega Norte")
    h(d, "1. Presentación")
    para(d, "Este manual explica cómo utilizar el MVP del sistema de inventario y ventas. Está dirigido al administrador o encargado de una bodega y usa un lenguaje operativo, con pasos breves para las tareas frecuentes.")
    h(d, "2. Ingreso y navegación")
    numbered(d, ["Abra la aplicación; en modo compartido confirme su cuenta y solicite habilitación de rol al administrador. En demo se simulan roles.", "Revise el Dashboard para conocer ventas del día, ingresos, productos y alertas.", "Use el menú lateral en escritorio o el botón Abrir menú en móvil.", "Seleccione Inventario, Productos, Nueva venta, Reportes o Configuración según la tarea."])
    h(d, "3. Registrar una venta")
    numbered(d, ["Seleccione Nueva venta.", "Busque un producto por nombre o categoría.", "Agregue el producto al carrito e indique la cantidad.", "Verifique subtotal y total; no ingrese una cantidad superior al stock disponible.", "Seleccione el método de pago: efectivo, Yape, Plin o tarjeta.", "Pulse Cobrar venta y luego Registrar venta.", "Confirme el mensaje de operación registrada y revise que el stock haya disminuido."])
    h(d, "4. Consultar inventario")
    numbered(d, ["Abra Inventario.", "Use el buscador para localizar un producto.", "Use el filtro Todos, Disponible o Bajo.", "Considere que un producto está bajo cuando su stock actual es menor o igual a su stock mínimo.", "Planifique la reposición tomando como referencia el listado y el reporte."])
    h(d, "5. Agregar un producto")
    numbered(d, ["Abra Productos y pulse Agregar producto.", "Complete nombre, categoría, precio, stock y stock mínimo.", "Revise que el precio sea mayor que cero y que las cantidades sean válidas.", "Pulse Guardar producto; se validan campos y valores. Para cambiar stock de un producto existente use Ajustar stock con motivo."])
    h(d, "6. Revisar reportes")
    bullets(d, ["En Reportes seleccione Desde y Hasta; revise ventas activas, ranking y medios de pago. Movimientos conserva ajustes y cancelaciones.", "Use Exportar CSV para descargar los datos y abrirlos en Excel.", "Conserve el archivo exportado como respaldo y no comparta datos sensibles sin autorización."])
    h(d, "7. Usar el asistente de bodega")
    numbered(d, ["Pulse el botón Asistente de bodega.", "Seleccione una consulta rápida como Stock bajo o Ventas de hoy.", "Lea la respuesta resumida y valide el detalle en Inventario o Reportes.", "La versión actual es una demostración offline; la integración productiva con Watson está documentada aparte."])
    h(d, "8. Recomendaciones de uso")
    bullets(d, ["Registre la venta inmediatamente después de atender al cliente.", "Revise el stock bajo al inicio y al cierre de la jornada.", "Si un cobro queda pendiente, reintente el mismo cobro; su identificador se conserva tras recargar y evita duplicarlo."])
    h(d, "9. Errores frecuentes")
    table(d, ["Situación", "Acción"], [
        ("No puedo guardar un producto", "Complete todos los campos y use valores válidos."),
        ("El producto no aparece", "Revise el buscador, categoría y que el catálogo esté cargado."),
        ("No puedo vender la cantidad", "La cantidad supera el stock disponible; reduzca unidades o reponga."),
        ("El reporte parece vacío", "Verifique que existan ventas registradas en el periodo de demostración."),
        ("Necesito conservar los datos", "El administrador descarga JSON en Configuración → Respaldo de datos; la recuperación requiere validación técnica independiente."),
    ])

    h(d, "9.1 Historial, ajustes y anulación", 2)
    numbered(d, ["Historial de ventas → Detalle permite revisar líneas e imprimir una constancia interna sin validez tributaria.", "El administrador usa Anular, registra el motivo y comprueba la devolución de stock. Repetir una anulación no devuelve las unidades dos veces.", "Para reponer o corregir existencias, use Ajustar stock en Inventario e indique cantidad positiva o negativa y motivo; editar el producto no modifica stock."])
    h(d, "9.2 Proveedores, permisos y respaldo", 2)
    numbered(d, ["Proveedores permite crear y editar contactos. Reasigne productos activos antes de eliminar un proveedor.", "Configuración → Usuarios y permisos permite habilitar cuentas confirmadas como cajero o administrador. El cajero consulta y vende; el administrador gestiona catálogo, ajustes, anulaciones y accesos.", "Configuración → Respaldo de datos descarga JSON sin contraseñas. Auditoría conserva operaciones y responsables. La restauración requiere un entorno de recuperación y validación técnica."])
    h(d, "Actualización operativa al 04-10-2026")
    para(d, "Se implementaron Supabase Auth y permisos administrador/cajero, ajustes con motivo y control de versión, venta con identificador de reintento, anulaciones con devolución única de stock, historial y constancia interna, proveedores y configuración persistentes, auditoría y respaldo JSON. Los reportes usan fecha de Lima y detalles de venta; no se inventan costos históricos.")
    para(d, "Validación técnica: 18 pruebas locales aprobadas; frontend y .NET compilan, con cero errores y advertencias en .NET; Java compila. La revisión de navegador se realizó en modo demo. La conexión privada del backend, integración HTTP PostgreSQL, despliegue Render y aceptación del Grupo 10 están pendientes. La evidencia detallada está en docs/evidencias/validacion_tecnica_2026-10-04.md.")
    save(d, "05_Manual_Usuario.docx")


def doc_06():
    d = Document(); cover(d, "Guía de instalación y puesta en marcha", "Documento operativo para el entorno local")
    h(d, "1. Alcance")
    para(d, "La guía permite ejecutar el MVP de Bodega Norte en un equipo de desarrollo. La solución incluye una View React alojada por ASP.NET Core MVC, controladores, modelos, reglas de negocio y repositorio Npgsql para Supabase PostgreSQL. También conserva un modo demo con persistencia local del navegador.")
    h(d, "2. Requisitos")
    bullets(d, ["Windows con Visual Studio, desarrollo ASP.NET y SDK de .NET 8.", "Node.js 22 y pnpm 11.19.0.", "Navegador actualizado: Chrome, Edge o Firefox.", "Proyecto Supabase PostgreSQL y contraseña de base de datos configurada en User Secrets.", "Java 17 o superior solo para ejecutar la prueba de integración opcional en una base de desarrollo."])
    h(d, "3. Instalación del frontend")
    numbered(d, ["Copie o clone la carpeta del proyecto en el equipo.", "Abra una terminal en la carpeta EFSR14908_Bodega_Norte.", "Ejecute npm install para instalar las dependencias.", "Ejecute npm run dev para iniciar el servidor local.", "Abra la URL indicada por Vite, normalmente http://localhost:5173/."])
    h(d, "4. Construcción de producción")
    para(d, "Para validar que la aplicación puede compilarse, ejecute:")
    p = d.add_paragraph(); r = p.add_run("npm run build"); r.font.name = "Consolas"; r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(BLUE)
    para(d, "El comando debe finalizar sin errores y crear la carpeta dist. Para preparar la aplicación MVC en Visual Studio, use npm run build:visualstudio; los archivos compilados se generan en backend/BodegaNorte.Api/wwwroot.")
    h(d, "4.1. Ejecución en Visual Studio y Supabase")
    numbered(d, ["Para una base nueva ejecute database/schema.sql, database/upgrade_operations.sql y database/seed.sql, en ese orden. En el proyecto actual la migración ya está aplicada.", "En Visual Studio, abra Manage User Secrets para BodegaNorte.Api y agregue ConnectionStrings:BodegaNorte con los datos del Session pooler y la contraseña, Supabase:Url, Supabase:PublishableKey y BootstrapAdminEmail. Confirme el correo del administrador inicial.", "Desde la raíz ejecute npm install y npm run build:visualstudio.", "Abra backend/BodegaNorte.sln en Visual Studio.", "Seleccione BodegaNorte.Api y presione F5.", "Verifique /api/health/ready; confirme su cuenta e ingrese con un rol habilitado para consultar productos."])
    para(d, "Formato de la conexión: Host=HOST_DEL_SESSION_POOLER;Port=5432;Database=postgres;Username=USUARIO_DEL_POOLER;Password=TU_CONTRASEÑA;SSL Mode=Require. Reemplace cada campo con los valores del proyecto Supabase; no copie la contraseña en el repositorio.")
    h(d, "5. Datos y persistencia")
    para(d, "El modo demo guarda datos en bodega-norte:v2:data y migra datos antiguos sin inventar sus fechas o costos. En el modo MVC, los controladores envían las operaciones a BodegaService y BodegaRepository las guarda en Supabase PostgreSQL dentro de transacciones.")
    table(d, ["Fase", "Persistencia", "Acción requerida"], [
        ("MVP demostrativo", "localStorage", "Usar Exportar CSV para respaldar."),
        ("MVC con Visual Studio", "Supabase PostgreSQL", "Aplicar schema.sql y seed.sql, configurar User Secrets y abrir la solución."),
        ("Producción", "Supabase PostgreSQL", "Validar roles y auditoría, guardar la conexión privada y ensayar recuperación antes de operar."),
    ])
    h(d, "6. Configuración ambiental")
    para(d, "Copie .env.example como .env solo si se agregan servicios externos. Nunca coloque claves reales en el repositorio ni en las capturas de la entrega.")
    table(d, ["Variable sugerida", "Uso"], [
        ("VITE_API_MODE", "El valor supabase activa la View conectada a MVC; demo activa LocalStorage."),
        ("VITE_API_BASE_URL", "Ruta base de los controladores MVC; por defecto /api."),
        ("WATSON_ASSISTANT_ID", "Identificador del asistente."),
        ("WATSON_APIKEY", "Clave privada; usar solo en servidor."),
        ("WATSON_SERVICE_URL", "Endpoint del servicio conversacional."),
    ])
    h(d, "7. Ejecución de prueba Java")
    para(d, "La prueba de integración está en integration/JavaInventoryIntegrationTest.java. Compile y ejecútela según integration/README.md. El resultado esperado confirma venta registrada y stock actualizado.")
    h(d, "8. Solución de problemas")
    table(d, ["Problema", "Solución"], [
        ("npm no reconocido", "Instale Node.js y confirme que npm quedó disponible en PATH."),
        ("Puerto ocupado", "Cierre otro proceso o inicie Vite con un puerto disponible."),
        ("La demo no guarda o muestra datos actualizados.", "Recargue y revise la consola. Exporte CSV como respaldo; localStorage no sustituye una base de datos."),
    ])

    h(d, "Actualización operativa al 04-10-2026")
    para(d, "Se implementaron Supabase Auth y permisos administrador/cajero, ajustes con motivo y control de versión, venta con identificador de reintento, anulaciones con devolución única de stock, historial y constancia interna, proveedores y configuración persistentes, auditoría y respaldo JSON. Los reportes usan fecha de Lima y detalles de venta; no se inventan costos históricos.")
    para(d, "Validación técnica: 18 pruebas locales aprobadas; frontend y .NET compilan, con cero errores y advertencias en .NET; Java compila. La revisión de navegador se realizó en modo demo. La conexión privada del backend, integración HTTP PostgreSQL, despliegue Render y aceptación del Grupo 10 están pendientes. La evidencia detallada está en docs/evidencias/validacion_tecnica_2026-10-04.md.")
    save(d, "06_Guia_Instalacion.docx")


def doc_07():
    d = Document(); cover(d, "Manual técnico", "Arquitectura, componentes y evolución de Bodega Norte")
    h(d, "1. Arquitectura actual")
    para(d, "El sistema se entrega con una arquitectura MVC estricta de tres capas. La View React se aloja en Views/Home/Index.cshtml; los Controllers reciben las solicitudes; el Model reúne contratos, reglas de negocio y persistencia en Supabase PostgreSQL mediante Npgsql. El modo demo usa localStorage únicamente como alternativa sin servidor.")
    table(d, ["Capa MVC", "Implementación", "Responsabilidad"], [
        ("View / Vista", "React + Vite + Views/Home/Index.cshtml", "Presentar módulos, formularios, navegación responsive y estados visuales."),
        ("Controller / Controlador", "Controllers/ + BodegaService.cs", "Recibir solicitudes, validar casos de uso y devolver vistas o JSON."),
        ("Model / Modelo", "Models/Contracts.cs + Data/BodegaRepository.cs", "Representar entidades, ejecutar consultas y persistir transacciones en Supabase PostgreSQL."),
        ("Demo alternativa", "localStorage", "Permitir una demostración sin servidor, sin formar parte del modo MVC."),
        ("Integración posterior", "Watson Assistant", "Consultas conversacionales mediante un controlador seguro."),
    ])
    h(d, "2. Estructura del proyecto")
    table(d, ["Ruta", "Contenido"], [
        ("src/main.jsx", "View React: composición visual, interacción y navegación."),
        ("src/styles.css", "Sistema visual, layout, responsive y estados."),
        ("backend/BodegaNorte.Api/Controllers/", "Controllers MVC para Home, productos, ventas y health."),
        ("backend/BodegaNorte.Api/Models/", "Contratos de entrada y salida del Model."),
        ("backend/BodegaNorte.Api/Services/", "Reglas de negocio de Bodega Norte."),
        ("backend/BodegaNorte.Api/Data/", "Repositorio Npgsql y transacciones PostgreSQL."),
        ("backend/BodegaNorte.Api/Views/", "Vista host MVC para la aplicación React."),
        ("database/schema.sql", "Tablas, relaciones, índices y restricciones."),
        ("integration/", "Prueba de integración Java y guía asociada."),
        ("docs/diagramas/", "Arquitectura draw.io y modelo de datos SVG."),
        ("docs/word/", "Informes y manuales editables para la entrega académica."),
    ])
    h(d, "3. Modelo de datos previsto")
    table(d, ["Entidad", "Campos principales", "Relación"], [
        ("categories", "category_id, name, is_active", "Una categoría agrupa productos."),
        ("suppliers", "supplier_id, name, contact_name, phone", "Un proveedor se asocia a productos."),
        ("products", "product_id, code, name, category_id, supplier_id, stock, min_stock, cost, sale_price", "Un producto pertenece a categoría y puede referenciar proveedor."),
        ("sales", "sale_id, sale_number, payment_method, total, sold_at", "Una venta contiene líneas de detalle."),
        ("sale_items", "sale_item_id, sale_id, product_id, quantity, unit_price", "Relaciona ventas y productos."),
        ("members", "user_id, email, role, is_active", "Cuenta de Supabase Auth habilitada para la bodega."),
        ("settings", "id, name, ruc, address, phone, stock_alerts", "Configuración única de la bodega."),
        ("stock_movements", "movement_id, product_id, quantity, balance, type, reason, actor", "Trazabilidad de altas, ventas, ajustes y anulaciones."),
        ("audit_log", "audit_id, action, entity, actor, details, created_at", "Registro de operaciones administrativas y de venta."),
    ])
    h(d, "4. Reglas de negocio")
    bullets(d, ["Una venta no debe superar el stock disponible.", "El total de una venta es la suma de cantidad por precio unitario.", "Stock bajo cuando stock actual es menor o igual a stock mínimo.", "El registro de una venta descuenta unidades y guarda el método de pago.", "Los datos de entrada deben validarse antes de persistir.", "Las claves de Watson deben permanecer fuera del cliente y del repositorio."])
    h(d, "5. Ejecución y evolución MVC")
    numbered(d, ["Crear el esquema Supabase con database/schema.sql, database/upgrade_operations.sql y database/seed.sql.", "Ejecutar npm run build:visualstudio para generar la View React en wwwroot.", "Abrir backend/BodegaNorte.sln y ejecutar BodegaNorte.Api desde Visual Studio.", "Mantener las reglas de stock y el registro de venta en BodegaService/Repository dentro de una transacción.", "Validar autenticación, roles, auditoría y concurrencia implementados en Controllers y Services.", "Extender los Models y Controllers para nuevos módulos sin mezclar SQL con la View."])
    h(d, "6. Integración Watson Assistant")
    para(d, "La integración documentada separa intents, entidades, respuestas y contrato de backend. El navegador no debe llamar directamente al servicio con una API key; el backend debe autenticar, filtrar la consulta y devolver solo información permitida para el usuario.")
    h(d, "7. Calidad y seguridad")
    bullets(d, ["Ejecutar npm run build antes de cada entrega.", "Mantener pruebas del flujo de venta y actualización de stock.", "Revisar accesibilidad de botones, formularios y navegación móvil.", "No subir .env, tokens, credenciales ni datos reales de clientes.", "Respaldar documentación y exportaciones fuera del navegador."])
    h(d, "8. Evidencia visual")
    image(d, "docs/evidencias/EVID_visual_escritorio.png", "Referencia visual anterior del dashboard", 6.5)

    h(d, "Actualización operativa al 04-10-2026")
    para(d, "Se implementaron Supabase Auth y permisos administrador/cajero, ajustes con motivo y control de versión, venta con identificador de reintento, anulaciones con devolución única de stock, historial y constancia interna, proveedores y configuración persistentes, auditoría y respaldo JSON. Los reportes usan fecha de Lima y detalles de venta; no se inventan costos históricos.")
    para(d, "Validación técnica: 18 pruebas locales aprobadas; frontend y .NET compilan, con cero errores y advertencias en .NET; Java compila. La revisión de navegador se realizó en modo demo. La conexión privada del backend, integración HTTP PostgreSQL, despliegue Render y aceptación del Grupo 10 están pendientes. La evidencia detallada está en docs/evidencias/validacion_tecnica_2026-10-04.md.")
    save(d, "07_Manual_Tecnico.docx")


def doc_08():
    d = Document(); cover(d, "Matriz de pruebas y validación", "Evidencias funcionales, técnicas y de usabilidad")
    h(d, "1. Alcance de la matriz")
    para(d, "La matriz separa las compilaciones realizadas de las pruebas de integración con Supabase y de la aceptación visual. La prueba SQL Server anterior no valida PostgreSQL. El 04-10-2026 pasaron 18 pruebas locales y se revisaron ventas, cancelación, proveedores, configuración, roles demo y fechas en el navegador. La integración HTTP PostgreSQL y la aceptación del equipo permanecen pendientes.")
    h(d, "2. Resultados técnicos")
    table(d, ["ID", "Acción", "Criterio", "Resultado y evidencia"], [
        ("BL-01", "npm run build", "Vite compila la interfaz sin errores.", "Aprobada el 28-09-2026; 16 módulos compilados."),
        ("BL-02", "dotnet build backend/BodegaNorte.sln", "La solución backend compila.", "Aprobada el 28-09-2026; cero advertencias y cero errores."),
        ("IT-API-01", "Registrar por API una venta de una unidad de P-001 en Supabase de desarrollo.", "Stock disminuye exactamente una unidad.", "Pendiente: requiere proyecto, conexión y prueba sobre base de desarrollo."),
        ("IT-API-02", "Solicitar por API una venta superior al stock en Supabase de desarrollo.", "HTTP 400 y stock sin cambios.", "Pendiente: requiere proyecto, conexión y prueba sobre base de desarrollo."),
    ])
    h(d, "3. Aceptación de interfaz pendiente")
    table(d, ["ID", "Acción", "Resultado esperado", "Estado"], [
        ("UI-01", "Abrir Dashboard.", "Indicadores, actividad y alertas visibles.", "Pendiente de revisión del equipo."),
        ("UI-02", "Crear y editar un producto.", "Cambios visibles en catálogo e inventario.", "Pendiente de revisión del equipo."),
        ("UI-03", "Aplicar filtro Bajo.", "Solo productos con stock menor o igual al mínimo.", "Pendiente de revisión del equipo."),
        ("UI-04", "Registrar venta desde el carrito.", "Confirmación y actualización de existencias.", "Pendiente de revisión del equipo."),
        ("UI-05", "Intentar seleccionar más unidades que el stock.", "La interfaz impide confirmar la venta.", "Pendiente de revisión del equipo."),
        ("UI-06", "Exportar CSV.", "Archivo válido y legible en hoja de cálculo.", "Pendiente de revisión del equipo."),
        ("UI-07", "Revisar ventana de 390 px.", "Controles utilizables sin desbordamiento crítico.", "Pendiente de revisión del equipo."),
        ("UI-08", "Consultar el asistente local.", "Respuesta de demostración; Watson no conectado.", "Pendiente de revisión del equipo."),
    ])
    h(d, "4. Registro de ejecución técnica")
    table(d, ["Fecha", "Ejecutor y ambiente", "Observación"], [
        ("28-09-2026", "Verificación técnica local para preparar el paquete; Windows, .NET SDK, Java 22 y SQL Server LocalDB en la versión previa.", "La prueba histórica creó una venta en una base desechable y luego se eliminó. No valida el backend Npgsql actual ni representa aceptación firmada del Grupo 10."),
    ])
    h(d, "5. Registro del equipo")
    para(d, "La aceptación manual debe ser registrada por integrantes del Grupo 10. No se incluyen firmas ni nombres de validadores que no hayan sido proporcionados.")
    table(d, ["Responsable", "Fecha", "Evidencia de aceptación"], [("Pendiente de registro por el equipo", "Pendiente", "Captura o acta por adjuntar")])

    h(d, "Actualización operativa al 04-10-2026")
    para(d, "Se implementaron Supabase Auth y permisos administrador/cajero, ajustes con motivo y control de versión, venta con identificador de reintento, anulaciones con devolución única de stock, historial y constancia interna, proveedores y configuración persistentes, auditoría y respaldo JSON. Los reportes usan fecha de Lima y detalles de venta; no se inventan costos históricos.")
    para(d, "Validación técnica: 18 pruebas locales aprobadas; frontend y .NET compilan, con cero errores y advertencias en .NET; Java compila. La revisión de navegador se realizó en modo demo. La conexión privada del backend, integración HTTP PostgreSQL, despliegue Render y aceptación del Grupo 10 están pendientes. La evidencia detallada está en docs/evidencias/validacion_tecnica_2026-10-04.md.")
    save(d, "08_Matriz_Pruebas.docx")


def doc_09():
    d = Document(); cover(d, "Guion de sustentación", "Defensa grupal del proyecto Bodega Norte")
    h(d, "1. Objetivo de la exposición")
    para(d, "Presentar en forma clara el problema, la solución, el proceso de desarrollo, las evidencias y la viabilidad del sistema. La exposición debe mostrar el producto funcionando y explicar qué se entrega en cada hito.")
    h(d, "2. Distribución sugerida")
    table(d, ["Bloque", "Tiempo sugerido", "Responsable", "Contenido"], [
        ("Contexto y propuesta", "4 min", COORDINATOR, "Problema del caso, objetivos, alcance y Canvas."),
        ("Diseño de datos", "3 min", TEAM[1], "Entidades SQL, relaciones, carga inicial y configuración."),
        ("Demostración", "5 min", TEAM[0], "Catálogo, venta, inventario y reportes."),
        ("Backend y pruebas", "3 min", TEAM[2], "API, compilación, integración y límites del prototipo."),
        ("Cierre y preguntas", "2 min", "Todo el equipo", "Conclusiones, mejoras y preguntas."),
    ])
    h(d, "3. Guion de demostración")
    numbered(d, ["Mostrar el Dashboard y explicar los cuatro indicadores principales.", "Abrir Inventario y aplicar el filtro Bajo para mostrar las alertas.", "Abrir Nueva venta, seleccionar Leche Gloria 1L y Gaseosa Inca Kola, elegir Yape y registrar la venta.", "Volver al inventario para mostrar que el stock se actualizó.", "Abrir Reportes y explicar el resumen y la exportación CSV.", "Abrir Asistente de bodega y consultar Stock bajo; aclarar que la conexión Watson queda preparada para la siguiente fase."])
    h(d, "4. Mensajes clave")
    bullets(d, ["Bodega Norte es un caso de estudio ficticio y utiliza datos de demostración.", "La versión SQL Server anterior comprobó el registro de una venta, el descuento de una unidad y el rechazo de una cantidad superior al stock; esta evidencia no valida la migración a PostgreSQL.", "La conexión Supabase, la prueba de API sobre PostgreSQL, la aceptación visual, la validación con usuarios reales y el impacto operativo todavía requieren evidencia.", "El modo productivo requiere autenticación, roles, auditoría, respaldo y despliegue seguro; Watson Assistant está especificado, pero no conectado al prototipo."])
    h(d, "5. Preguntas previsibles")
    table(d, ["Pregunta", "Respuesta sugerida"], [
        ("¿Por qué localStorage?", "Es una alternativa demo sin servidor; la ejecución en Visual Studio usa MVC y Supabase PostgreSQL."),
        ("¿Cómo se evita vender sin stock?", "La cantidad se valida contra el stock disponible antes de confirmar."),
        ("¿Qué hace el asistente?", "La demo responde consultas de stock; el contrato Watson define la integración segura futura."),
        ("¿Cómo se respalda la información?", "La demo permite exportar CSV; la base Supabase requiere definir y comprobar una política de copias antes de producción."),
        ("¿Cómo se probó?", "Se compilaron frontend y backend MVC con Npgsql. La integración Java validó una versión anterior con SQL Server; la prueba HTTP equivalente en Supabase permanece pendiente. Las 18 pruebas locales y los flujos de navegador de demostración se verificaron el 04-10-2026."),
    ])
    h(d, "6. Cierre sugerido")
    para(d, "La demostración presenta el alcance técnico del prototipo y diferencia las funciones implementadas de las validaciones pendientes. Antes de proponer un despliegue se deben confirmar la fecha del curso, la aceptación del equipo y los requisitos de un usuario real.")
    h(d, "7. Datos para completar")
    table(d, ["Dato", "Referencia"], [("Integrantes", "Grupo 10; relación en la portada."), ("Orden de exposición", "Propuesta en la sección 2; el equipo puede ajustarla."), ("Duración", "17 minutos sugeridos, sujeto a indicación de la docente."), ("Fecha y aula", "Según anuncio vigente en el aula virtual.")])

    h(d, "Actualización operativa al 04-10-2026")
    para(d, "Se implementaron Supabase Auth y permisos administrador/cajero, ajustes con motivo y control de versión, venta con identificador de reintento, anulaciones con devolución única de stock, historial y constancia interna, proveedores y configuración persistentes, auditoría y respaldo JSON. Los reportes usan fecha de Lima y detalles de venta; no se inventan costos históricos.")
    para(d, "Validación técnica: 18 pruebas locales aprobadas; frontend y .NET compilan, con cero errores y advertencias en .NET; Java compila. La revisión de navegador se realizó en modo demo. La conexión privada del backend, integración HTTP PostgreSQL, despliegue Render y aceptación del Grupo 10 están pendientes. La evidencia detallada está en docs/evidencias/validacion_tecnica_2026-10-04.md.")
    save(d, "09_Guion_Sustentacion.docx")


def doc_10():
    d = Document(); cover(d, "Especificación de integración con Watson Assistant", "Contrato funcional y técnico para una fase posterior")
    h(d, "1. Propósito")
    para(d, "Este documento define cómo incorporar un asistente conversacional al sistema Bodega Norte. La versión actual incluye un asistente offline demostrativo; este diseño permite conectar Watson Assistant sin exponer credenciales ni reglas internas en el navegador.")
    h(d, "2. Intents propuestos")
    table(d, ["Intent", "Ejemplos", "Respuesta esperada"], [
        ("stock_bajo", "¿Qué productos están bajos?", "Lista de productos cuyo stock actual está en o debajo del mínimo."),
        ("ventas_hoy", "¿Cuánto vendimos hoy?", "Total de ventas e ingresos del día."),
        ("buscar_producto", "Busca arroz", "Producto, precio y disponibilidad."),
        ("resumen_inventario", "Dame un resumen del inventario", "Cantidad de productos, valorización referencial y alertas."),
        ("ayuda_sistema", "¿Cómo registro una venta?", "Pasos resumidos para completar la tarea."),
    ])
    h(d, "3. Entidades")
    table(d, ["Entidad", "Valores / ejemplo", "Uso"], [
        ("producto", "Leche Gloria 1L", "Buscar un artículo específico."),
        ("categoria", "Lácteos, bebidas, abarrotes", "Filtrar catálogo."),
        ("periodo", "hoy, semana, mes", "Delimitar reportes."),
        ("metodo_pago", "efectivo, Yape, Plin, tarjeta", "Consultar ventas por forma de pago."),
    ])
    h(d, "4. Rutas propuestas para una integración futura")
    para(d, "Las rutas de esta tabla son una propuesta y no están implementadas en la API actual.")
    table(d, ["Método", "Ruta", "Propósito"], [
        ("GET", "/api/inventory/low-stock", "Obtener productos con stock bajo."),
        ("GET", "/api/reports/sales?period=today", "Obtener resumen de ventas."),
        ("GET", "/api/products/search?q=...", "Buscar producto."),
        ("GET", "/api/inventory/summary", "Obtener resumen de inventario."),
    ])
    h(d, "5. Flujo seguro")
    numbered(d, ["El usuario envía una pregunta desde la interfaz.", "El frontend envía la consulta al backend de Bodega Norte.", "El backend valida usuario, permisos y parámetros.", "El backend llama a Watson Assistant usando credenciales protegidas.", "El backend consulta solo los datos autorizados y forma una respuesta segura.", "El frontend presenta la respuesta y ofrece un enlace al módulo correspondiente."])
    h(d, "6. Reglas de seguridad")
    bullets(d, ["Nunca colocar WATSON_APIKEY en el código cliente ni en una captura.", "Usar variables de entorno en el backend y rotar credenciales.", "Limitar respuestas a datos necesarios para el rol del usuario.", "Registrar errores sin guardar preguntas que contengan información sensible.", "Aplicar rate limiting y validar entradas antes de consultar la base de datos."])
    h(d, "7. Criterios de aceptación")
    table(d, ["Criterio", "Evidencia"], [
        ("El intent stock_bajo reconoce variaciones de la pregunta.", "Pruebas conversacionales documentadas."),
        ("La respuesta coincide con el inventario autorizado.", "Comparación con /api/inventory/low-stock."),
        ("No se exponen credenciales.", "Revisión de variables y repositorio."),
        ("El asistente funciona sin romper la venta.", "Prueba de regresión IT-02."),
    ])
    h(d, "8. Estado actual")
    para(d, "Estado: especificación preparada; integración productiva pendiente de credenciales, configuración de Watson, backend y validación institucional. La aplicación conserva una respuesta offline para que la demostración del proyecto sea autónoma.")
    save(d, "10_Integracion_Watson_Assistant.docx")


def main():
    # Regenera los entregables del proyecto sin borrar archivos ajenos.
    doc_01(); doc_02(); doc_03(); doc_04(); doc_05(); doc_06(); doc_07(); doc_08(); doc_09(); doc_10()
    print(f"Created {len(list(OUT.glob('*.docx')))} Word documents in {OUT}")


if __name__ == "__main__":
    main()
